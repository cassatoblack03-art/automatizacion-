#!/usr/bin/env python3
"""Genera el proyecto Power BI (.pbip) del reporte ejecutivo
"Backorder en CMMF viejo -> cambio al CMMF nuevo" a partir del Excel de corte.

Uso:
    python generar_pbip.py <archivo.xlsx> <carpeta_salida>

El proyecto resultante trae los datos del corte embebidos (abre sin depender
del Excel) y un parámetro `RutaExcel` para actualizarlo con un corte nuevo
desde Power BI Desktop.
"""

import base64
import datetime
import json
import os
import re
import shutil
import sys
import uuid
import zlib

import openpyxl

NOMBRE = "Reporte BO CMMF"
NS = uuid.UUID("6f1c2b8e-3d4a-4c4e-9a51-0b8e2f7d9c11")

# Paleta corporativa (skill dashboard-ejecutivo-rojo)
ROJO = "#B30000"
ROJO_OSCURO = "#7D0000"
ROJO_SUAVE = "#FDEAEA"
ROJO_MEDIO = "#E06666"
FONDO = "#F8F9FA"
BLANCO = "#FFFFFF"
BORDE = "#E6E8EB"
TEXTO = "#212529"
TEXTO_2 = "#6C757D"
GRIS = "#ADB5BD"

FACT = "BO Lineas"
DIM = "Productos CMMF"
MED = "Medidas"
PUNTOS = "Puntos Clave"

SCHEMA = "https://developer.microsoft.com/json-schemas/fabric/"
VC_SCHEMA = SCHEMA + "item/report/definition/visualContainer/1.3.0/schema.json"
BASE_THEME = "CY24SU10"
THEME_VERSION = "5.59"
CUSTOM_THEME = "BO_Ejecutivo_Rojo.json"


def gid(key):
    return str(uuid.uuid5(NS, key))


def hid(key):
    return uuid.uuid5(NS, key).hex[:20]


# ---------------------------------------------------------------------------
# 1. Lectura del Excel
# ---------------------------------------------------------------------------

def leer_hoja(ws, primera_columna):
    filas = [[c.value for c in r] for r in ws.iter_rows()]
    inicio = next(i for i, r in enumerate(filas) if str(r[0]).strip().upper() == primera_columna)
    encabezados = [str(h).strip() for h in filas[inicio] if h is not None]
    datos = []
    for r in filas[inicio + 1:]:
        if r[0] is None or str(r[0]).strip().upper() == "TOTAL":
            continue
        datos.append([valor_json(v) for v in r[:len(encabezados)]])
    return encabezados, datos


def valor_json(v):
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return int(v)
    return v


def embebido(filas):
    """Codifica las filas igual que 'Especificar datos' de Power BI (deflate crudo + base64)."""
    crudo = json.dumps(filas, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    comp = zlib.compressobj(9, zlib.DEFLATED, -15)
    datos = comp.compress(crudo) + comp.flush()
    assert json.loads(zlib.decompress(datos, -15).decode("utf-8")) == filas
    return base64.b64encode(datos).decode("ascii")


# ---------------------------------------------------------------------------
# 2. Modelo semántico (TMDL)
# ---------------------------------------------------------------------------

def q(nombre):
    return nombre if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", nombre) else "'" + nombre.replace("'", "''") + "'"


def m_lista(valores):
    return "{" + ", ".join('"' + v.replace('"', '""') + '"' for v in valores) + "}"


FACT_RENOMBRES = [
    ("CANAL", "Canal"), ("GRUPO", "Grupo"), ("NOMBRE MLA", "Nombre MLA"), ("CODIGO 3A", "Cliente"),
    ("LINEA", "Linea"), ("FECHA PEDIDO", "Fecha Pedido"), ("DIAS", "Dias BO"),
    ("FECHA CANCELACION OC", "Fecha Cancelacion OC"), ("ESTADO OC", "Estado OC Origen"),
    ("CMMF VIEJO", "CMMF Viejo"), ("CODIGO COMERCIAL VIEJO", "Codigo Comercial Viejo"),
    ("DESCR VIEJO", "Producto Viejo"), ("PZAS BO", "Pzas BO"), ("PRECIO NETO", "Precio Neto"),
    ("$ BO", "Monto BO"), ("CUENTA COMO", "Cuenta Como"), ("PZAS PICKING VIEJO", "Pzas Picking Viejo"),
    ("PICKING ID", "Picking ID"), ("CMMF NUEVO", "CMMF Nuevo"),
    ("CODIGO COMERCIAL NUEVO", "Codigo Comercial Nuevo"), ("DESCR NUEVO", "Producto Nuevo"),
    ("OV YA TRAE EL NUEVO", "OV Ya Trae Nuevo"), ("PZAS BO DEL NUEVO EN LA OV", "Pzas BO Nuevo En OV"),
    ("QUE HACER", "Que Hacer"),
]

M_TIPO = {"string": "type text", "int64": "Int64.Type", "double": "type number", "dateTime": "type date"}

# (nombre, tipo TMDL, formato, resumen, oculto, ordenar por, carpeta)
FACT_COLUMNAS = [
    ("Canal", "string", None, "none", False, None, "Cliente"),
    ("Grupo", "string", None, "none", False, None, "Cliente"),
    ("MLA", "string", None, "none", False, None, "Cliente"),
    ("Nombre MLA", "string", None, "none", False, None, "Cliente"),
    ("Cliente", "string", None, "none", False, None, "Cliente"),
    ("Cliente Corto", "string", None, "none", False, None, "Cliente"),
    ("KAM", "string", None, "none", False, None, "Cliente"),
    ("OV", "string", None, "none", False, None, "Pedido"),
    ("OC", "string", None, "none", False, None, "Pedido"),
    ("Linea", "int64", "0", "none", False, None, "Pedido"),
    ("Linea ID", "string", None, "none", True, None, "Pedido"),
    ("Fecha Pedido", "dateTime", "dd/mm/yyyy", "none", False, None, "Pedido"),
    ("Dias BO", "int64", "0", "none", False, None, "Pedido"),
    ("Antiguedad BO", "string", None, "none", False, "Antiguedad Orden", "Pedido"),
    ("Antiguedad Orden", "int64", "0", "none", True, None, "Pedido"),
    ("Fecha Cancelacion OC", "dateTime", "dd/mm/yyyy", "none", False, None, "Pedido"),
    ("Dias Para Cancelar OC", "int64", "0", "none", False, None, "Pedido"),
    ("Estado OC", "string", None, "none", False, "Estado OC Orden", "Pedido"),
    ("Estado OC Orden", "int64", "0", "none", True, None, "Pedido"),
    ("Estado OC Origen", "string", None, "none", True, None, "Pedido"),
    ("Fecha Corte", "dateTime", "dd/mm/yyyy", "none", True, None, "Pedido"),
    ("CMMF Viejo", "string", None, "none", False, None, "Producto"),
    ("Codigo Comercial Viejo", "string", None, "none", False, None, "Producto"),
    ("Producto Viejo", "string", None, "none", False, None, "Producto"),
    ("CMMF Nuevo", "string", None, "none", False, None, "Producto"),
    ("Codigo Comercial Nuevo", "string", None, "none", False, None, "Producto"),
    ("Producto Nuevo", "string", None, "none", False, None, "Producto"),
    ("Cambio CMMF", "string", None, "none", False, None, "Producto"),
    ("Pzas BO", "int64", "#,0", "sum", False, None, "Cantidades"),
    ("Precio Neto", "double", "$#,0.00", "none", False, None, "Cantidades"),
    ("Monto BO", "double", "$#,0", "sum", False, None, "Cantidades"),
    ("Pzas Picking Viejo", "int64", "#,0", "sum", False, None, "Cantidades"),
    ("Pzas BO Nuevo En OV", "int64", "#,0", "sum", False, None, "Cantidades"),
    ("Cuenta Como", "string", None, "none", False, None, "Acción"),
    ("Picking ID", "string", None, "none", False, None, "Acción"),
    ("OV Ya Trae Nuevo", "string", None, "none", False, None, "Acción"),
    ("Que Hacer", "string", None, "none", False, None, "Acción"),
    ("Accion", "string", None, "none", False, "Accion Orden", "Acción"),
    ("Accion Orden", "int64", "0", "none", True, None, "Acción"),
]

DIM_COLUMNAS = [
    ("CMMF Viejo", "string", None, "none", False, None, None),
    ("Codigo Comercial Viejo", "string", None, "none", False, None, None),
    ("Producto Viejo", "string", None, "none", False, None, None),
    ("CMMF Nuevo", "string", None, "none", False, None, None),
    ("Codigo Comercial Nuevo", "string", None, "none", False, None, None),
    ("Producto Nuevo", "string", None, "none", False, None, None),
    ("Cambio CMMF", "string", None, "none", False, None, None),
]

PUNTOS_COLUMNAS = [
    ("Orden", "int64", "0", "none", True, None, None),
    ("Tema", "string", None, "none", False, "Orden", None),
]

PUNTOS_FILAS = [
    (1, "Cliente principal"),
    (2, "Producto principal"),
    (3, "Cambio inmediato"),
    (4, "OC vencida"),
    (5, "OC por vencer"),
    (6, "Revisar duplicados"),
]


def m_leer_excel(hoja, primera_columna):
    return f"""    LeerExcel = (ruta as text) as table =>
        let
            Libro = Excel.Workbook(File.Contents(ruta), null, true),
            Hoja = Libro{{[Item = "{hoja}", Kind = "Sheet"]}}[Data],
            Fila = List.PositionOf(List.Transform(Hoja[Column1], each Text.Upper(Text.Trim(Text.From(_)))), "{primera_columna}"),
            Datos = Table.PromoteHeaders(Table.Skip(Hoja, Fila), [PromoteAllScalars = true]),
            Seleccion = Table.SelectColumns(Datos, Columnas, MissingField.UseNull)
        in
            Seleccion,
    // Si RutaExcel está vacío se usan los datos del corte guardados en el archivo
    Origen = if RutaExcel = null or Text.Trim(RutaExcel) = "" then Embebido else LeerExcel(Text.Trim(RutaExcel)),"""


def m_fact(columnas_excel, b64):
    renombres = ",\n        ".join(f'{{"{a}", "{b}"}}' for a, b in FACT_RENOMBRES)
    base = {b for _, b in FACT_RENOMBRES} | {"MLA", "KAM", "OV", "OC"}
    tipos = ",\n        ".join(
        f'{{"{c[0]}", {M_TIPO[c[1]]}}}' for c in FACT_COLUMNAS if c[0] in base)
    texto = [c[0] for c in FACT_COLUMNAS if c[0] in base and c[1] == "string"]
    limpieza = ",\n        ".join(f'{{"{c}", Text.Trim, type text}}' for c in texto)
    finales = m_lista([c[0] for c in FACT_COLUMNAS])
    return f"""let
    Columnas = {m_lista(columnas_excel)},
    Embebido = Table.FromRows(Json.Document(Binary.Decompress(Binary.FromText("{b64}", BinaryEncoding.Base64), Compression.Deflate)), Columnas),
{m_leer_excel("Por linea SAP", "CANAL")}
    SoloLineas = Table.SelectRows(Origen, each [CANAL] <> null and Text.Upper(Text.Trim(Text.From([CANAL]))) <> "TOTAL" and [#"$ BO"] <> null),
    Renombrado = Table.RenameColumns(SoloLineas, {{
        {renombres}}}),
    Tipos = Table.TransformColumnTypes(Renombrado, {{
        {tipos}}}, "en-US"),
    Limpio = Table.TransformColumns(Tipos, {{
        {limpieza}}}),
    ConKam = Table.TransformColumns(Limpio, {{{{"KAM", each if _ = null or _ = "" then "SIN KAM ASIGNADO" else _, type text}}}}),
    ConClienteCorto = Table.AddColumn(ConKam, "Cliente Corto", each
        if [Cliente] = null then "SIN CLIENTE"
        else let p = Text.PositionOf([Cliente], "·") in if p >= 0 then Text.Trim(Text.Middle([Cliente], p + 1)) else [Cliente], type text),
    ConLineaId = Table.AddColumn(ConClienteCorto, "Linea ID", each Text.From([OV]) & "-" & Text.From([Linea]), type text),
    ConCorte = Table.AddColumn(ConLineaId, "Fecha Corte", each
        if [Fecha Pedido] = null or [Dias BO] = null then null else Date.AddDays([Fecha Pedido], [Dias BO]), type date),
    ConDiasCancelar = Table.AddColumn(ConCorte, "Dias Para Cancelar OC", each
        if [Fecha Cancelacion OC] = null or [Fecha Corte] = null then null else Duration.Days([Fecha Cancelacion OC] - [Fecha Corte]), Int64.Type),
    ConEstado = Table.AddColumn(ConDiasCancelar, "Estado OC", each
        let e = if [Estado OC Origen] = null then "" else Text.Upper([Estado OC Origen]) in
        if e = "VENCIDA" then "Vencida"
        else if Text.StartsWith(e, "VENCE") then "Vence en 7 días"
        else if e = "VIGENTE" then "Vigente"
        else "Sin estado", type text),
    ConEstadoOrden = Table.AddColumn(ConEstado, "Estado OC Orden", each
        if [Estado OC] = "Vencida" then 1 else if [Estado OC] = "Vence en 7 días" then 2 else if [Estado OC] = "Vigente" then 3 else 9, Int64.Type),
    ConAccion = Table.AddColumn(ConEstadoOrden, "Accion", each
        let a = if [Que Hacer] = null then "" else Text.Upper([Que Hacer]) in
        if Text.StartsWith(a, "CAMBIAR AL CMMF NUEVO") then "Cambiar al CMMF nuevo"
        else if Text.StartsWith(a, "CAMBIAR SOLO EL BO") then "Cambiar solo el BO"
        else if Text.StartsWith(a, "REVISAR") then "Revisar: OV ya trae el nuevo"
        else if Text.StartsWith(a, "CONFIRMAR") then "Confirmar con cliente"
        else "Otra acción", type text),
    ConAccionOrden = Table.AddColumn(ConAccion, "Accion Orden", each
        if [Accion] = "Cambiar al CMMF nuevo" then 1
        else if [Accion] = "Cambiar solo el BO" then 2
        else if [Accion] = "Revisar: OV ya trae el nuevo" then 3
        else if [Accion] = "Confirmar con cliente" then 4
        else 9, Int64.Type),
    ConAntiguedad = Table.AddColumn(ConAccionOrden, "Antiguedad BO", each
        if [Dias BO] = null then "Sin dato"
        else if [Dias BO] <= 7 then "0-7 días"
        else if [Dias BO] <= 15 then "8-15 días"
        else if [Dias BO] <= 30 then "16-30 días"
        else "Más de 30 días", type text),
    ConAntiguedadOrden = Table.AddColumn(ConAntiguedad, "Antiguedad Orden", each
        if [Dias BO] = null then 9
        else if [Dias BO] <= 7 then 1
        else if [Dias BO] <= 15 then 2
        else if [Dias BO] <= 30 then 3
        else 4, Int64.Type),
    ConCambio = Table.AddColumn(ConAntiguedadOrden, "Cambio CMMF", each
        (if [CMMF Viejo] = null then "" else [CMMF Viejo]) & " → " & (if [CMMF Nuevo] = null then "" else [CMMF Nuevo]), type text),
    Resultado = Table.SelectColumns(ConCambio, {finales})
in
    Resultado"""


def m_dim(columnas_excel, b64):
    return f"""let
    Columnas = {m_lista(columnas_excel)},
    Embebido = Table.FromRows(Json.Document(Binary.Decompress(Binary.FromText("{b64}", BinaryEncoding.Base64), Compression.Deflate)), Columnas),
{m_leer_excel("Por CMMF", "CMMF VIEJO")}
    Filtrado = Table.SelectRows(Origen, each [#"CMMF VIEJO"] <> null and Text.Upper(Text.Trim(Text.From([#"CMMF VIEJO"]))) <> "TOTAL" and Text.Trim(Text.From([#"CMMF VIEJO"])) <> ""),
    Renombrado = Table.RenameColumns(Filtrado, {{
        {{"CMMF VIEJO", "CMMF Viejo"}}, {{"CODIGO COMERCIAL VIEJO", "Codigo Comercial Viejo"}}, {{"DESCR VIEJO", "Producto Viejo"}},
        {{"CMMF NUEVO", "CMMF Nuevo"}}, {{"CODIGO COMERCIAL NUEVO", "Codigo Comercial Nuevo"}}, {{"DESCR NUEVO", "Producto Nuevo"}}}}),
    Tipos = Table.TransformColumnTypes(Renombrado, {{
        {{"CMMF Viejo", type text}}, {{"Codigo Comercial Viejo", type text}}, {{"Producto Viejo", type text}},
        {{"CMMF Nuevo", type text}}, {{"Codigo Comercial Nuevo", type text}}, {{"Producto Nuevo", type text}}}}, "en-US"),
    Limpio = Table.TransformColumns(Tipos, {{
        {{"CMMF Viejo", Text.Trim, type text}}, {{"Codigo Comercial Viejo", Text.Trim, type text}}, {{"Producto Viejo", Text.Trim, type text}},
        {{"CMMF Nuevo", Text.Trim, type text}}, {{"Codigo Comercial Nuevo", Text.Trim, type text}}, {{"Producto Nuevo", Text.Trim, type text}}}}),
    SinDuplicados = Table.Distinct(Limpio, {{"CMMF Viejo"}}),
    ConCambio = Table.AddColumn(SinDuplicados, "Cambio CMMF", each
        [CMMF Viejo] & " → " & (if [CMMF Nuevo] = null then "" else [CMMF Nuevo]), type text)
in
    ConCambio"""


def m_puntos():
    filas = ", ".join(f'{{{o}, "{t}"}}' for o, t in PUNTOS_FILAS)
    return f"""let
    Origen = #table(type table [Orden = Int64.Type, Tema = text], {{{filas}}})
in
    Origen"""


def m_medidas():
    return """let
    Origen = #table(type table [Medida = text], {})
in
    Origen"""


# Medidas: (nombre, expresión DAX, formato, carpeta)
MEDIDAS = [
    # 1. KPIs
    ("$ BO", "SUM('BO Lineas'[Monto BO])", "$#,0", "1. KPIs"),
    ("Piezas BO", "SUM('BO Lineas'[Pzas BO])", "#,0", "1. KPIs"),
    ("Líneas", "COUNTROWS('BO Lineas')", "#,0", "1. KPIs"),
    ("OVs", "DISTINCTCOUNT('BO Lineas'[OV])", "#,0", "1. KPIs"),
    ("Clientes", "DISTINCTCOUNT('BO Lineas'[Cliente])", "#,0", "1. KPIs"),
    ("CMMF con BO", "DISTINCTCOUNT('BO Lineas'[CMMF Viejo])", "#,0", "1. KPIs"),
    ("CMMF en Lista", "COUNTROWS(ALL('Productos CMMF'))", "#,0", "1. KPIs"),
    ("Pzas en Picking Viejo", "SUM('BO Lineas'[Pzas Picking Viejo])", "#,0", "1. KPIs"),
    ("$ BO OC Vencida", "CALCULATE([$ BO], KEEPFILTERS('BO Lineas'[Estado OC] = \"Vencida\"))", "$#,0", "1. KPIs"),
    ("% BO OC Vencida", "DIVIDE([$ BO OC Vencida], [$ BO])", "0.0%", "1. KPIs"),
    ("$ BO OC por Vencer", "CALCULATE([$ BO], KEEPFILTERS('BO Lineas'[Estado OC] = \"Vence en 7 días\"))", "$#,0", "1. KPIs"),
    ("Líneas OC por Vencer", "COALESCE(CALCULATE([Líneas], KEEPFILTERS('BO Lineas'[Estado OC] = \"Vence en 7 días\")), 0)", "#,0", "1. KPIs"),
    ("$ Cambio Inmediato", "CALCULATE([$ BO], KEEPFILTERS('BO Lineas'[Accion Orden] IN {1, 2}))", "$#,0", "1. KPIs"),
    ("Líneas Cambio Inmediato", "COALESCE(CALCULATE([Líneas], KEEPFILTERS('BO Lineas'[Accion Orden] IN {1, 2})), 0)", "#,0", "1. KPIs"),
    ("% del BO", "DIVIDE([$ BO], CALCULATE([$ BO], ALLSELECTED('BO Lineas')))", "0.0%", "1. KPIs"),
    ("Días Promedio en BO", "AVERAGE('BO Lineas'[Dias BO])", "#,0.0", "1. KPIs"),
    ("Fecha Corte", "CALCULATE(MAX('BO Lineas'[Fecha Corte]), REMOVEFILTERS('BO Lineas'))", "dd/mm/yyyy", "1. KPIs"),
    # 2. Por acción
    ("$ Acción Cambiar", "CALCULATE([$ BO], KEEPFILTERS('BO Lineas'[Accion Orden] = 1))", "$#,0", "2. Por acción"),
    ("$ Acción Solo BO", "CALCULATE([$ BO], KEEPFILTERS('BO Lineas'[Accion Orden] = 2))", "$#,0", "2. Por acción"),
    ("$ Acción Revisar", "CALCULATE([$ BO], KEEPFILTERS('BO Lineas'[Accion Orden] = 3))", "$#,0", "2. Por acción"),
    ("$ Acción Confirmar", "CALCULATE([$ BO], KEEPFILTERS('BO Lineas'[Accion Orden] = 4))", "$#,0", "2. Por acción"),
    ("Líneas Acción Cambiar", "CALCULATE([Líneas], KEEPFILTERS('BO Lineas'[Accion Orden] = 1))", "#,0", "2. Por acción"),
    ("Líneas Acción Solo BO", "CALCULATE([Líneas], KEEPFILTERS('BO Lineas'[Accion Orden] = 2))", "#,0", "2. Por acción"),
    ("Líneas Acción Revisar", "CALCULATE([Líneas], KEEPFILTERS('BO Lineas'[Accion Orden] = 3))", "#,0", "2. Por acción"),
    ("Líneas Acción Confirmar", "CALCULATE([Líneas], KEEPFILTERS('BO Lineas'[Accion Orden] = 4))", "#,0", "2. Por acción"),
    # 3. Textos dinámicos
    ("Texto Fecha Corte", """VAR _f = [Fecha Corte]
VAR _meses = "ene|feb|mar|abr|may|jun|jul|ago|sep|oct|nov|dic"
RETURN
    IF(ISBLANK(_f), "sin datos", FORMAT(DAY(_f), "00") & "-" & PATHITEM(_meses, MONTH(_f)) & "-" & YEAR(_f))""", None, "3. Textos"),
    ("Subtítulo Reporte", """"Corte " & [Texto Fecha Corte]
    & "   ·   " & FORMAT(COALESCE([Líneas], 0), "#,0") & " líneas"
    & "   ·   " & FORMAT(COALESCE([OVs], 0), "#,0") & " OVs"
    & "   ·   " & FORMAT(COALESCE([Clientes], 0), "#,0") & " clientes"
    & "   ·   " & FORMAT(COALESCE([CMMF con BO], 0), "#,0") & " CMMF viejos\"""", None, "3. Textos"),
    ("Sub $ BO", """"en " & FORMAT(COALESCE([Líneas], 0), "#,0") & " líneas · " & FORMAT(COALESCE([OVs], 0), "#,0") & " OVs\"""", None, "3. Textos"),
    ("Sub Piezas", """FORMAT(COALESCE([Pzas en Picking Viejo], 0), "#,0") & " pzas ya en picking (no se cambian)\"""", None, "3. Textos"),
    ("Sub Clientes", """FORMAT(COALESCE([CMMF con BO], 0), "#,0") & " de " & FORMAT([CMMF en Lista], "#,0") & " CMMF viejos con BO\"""", None, "3. Textos"),
    ("Sub OC Vencida", """FORMAT(COALESCE([$ BO OC Vencida], 0), "$#,0") & " por confirmar con el cliente\"""", None, "3. Textos"),
    ("Sub Cambio Inmediato", """FORMAT([Líneas Cambio Inmediato], "#,0") & " líneas listas para cambiar ya\"""", None, "3. Textos"),
    ("Sub OC por Vencer", """FORMAT([Líneas OC por Vencer], "#,0") & " líneas vencen en ≤7 días: priorizar\"""", None, "3. Textos"),
    ("Sub Acción Cambiar", """FORMAT(COALESCE([Líneas Acción Cambiar], 0), "#,0") & " líneas · " & FORMAT(COALESCE(DIVIDE([$ Acción Cambiar], [$ BO]), 0), "0%") & " del BO\"""", None, "3. Textos"),
    ("Sub Acción Solo BO", """FORMAT(COALESCE([Líneas Acción Solo BO], 0), "#,0") & " líneas · " & FORMAT(COALESCE(DIVIDE([$ Acción Solo BO], [$ BO]), 0), "0%") & " del BO\"""", None, "3. Textos"),
    ("Sub Acción Revisar", """FORMAT(COALESCE([Líneas Acción Revisar], 0), "#,0") & " líneas · " & FORMAT(COALESCE(DIVIDE([$ Acción Revisar], [$ BO]), 0), "0%") & " del BO\"""", None, "3. Textos"),
    ("Sub Acción Confirmar", """FORMAT(COALESCE([Líneas Acción Confirmar], 0), "#,0") & " líneas · " & FORMAT(COALESCE(DIVIDE([$ Acción Confirmar], [$ BO]), 0), "0%") & " del BO\"""", None, "3. Textos"),
    ("Punto Cliente", """VAR _total = [$ BO]
VAR _tabla = ADDCOLUMNS(VALUES('BO Lineas'[Cliente Corto]), "@monto", [$ BO])
VAR _top = TOPN(1, _tabla, [@monto], DESC)
VAR _cliente = MAXX(_top, 'BO Lineas'[Cliente Corto])
VAR _monto = MAXX(_top, [@monto])
VAR _vencido = CALCULATE([$ BO OC Vencida], 'BO Lineas'[Cliente Corto] = _cliente)
RETURN
    _cliente & " concentra el " & FORMAT(DIVIDE(_monto, _total), "0%") & " del BO (" & FORMAT(_monto, "$#,0") & ")"
        & IF(_vencido > 0, " y el " & FORMAT(DIVIDE(_vencido, _monto), "0%") & " de ese monto tiene la OC vencida: confirmar con el cliente.", ".")""", None, "3. Textos"),
    ("Punto Producto", """VAR _total = [$ BO]
VAR _tabla = ADDCOLUMNS(VALUES('BO Lineas'[CMMF Viejo]), "@monto", [$ BO])
VAR _top = TOPN(1, _tabla, [@monto], DESC)
VAR _cmmf = MAXX(_top, 'BO Lineas'[CMMF Viejo])
VAR _monto = MAXX(_top, [@monto])
VAR _producto = CALCULATE(MAX('BO Lineas'[Producto Viejo]), 'BO Lineas'[CMMF Viejo] = _cmmf)
VAR _nuevo = CALCULATE(MAX('BO Lineas'[CMMF Nuevo]), 'BO Lineas'[CMMF Viejo] = _cmmf)
RETURN
    _producto & " (CMMF " & _cmmf & " → " & _nuevo & ") representa el " & FORMAT(DIVIDE(_monto, _total), "0%") & " del BO (" & FORMAT(_monto, "$#,0") & ").\"""", None, "3. Textos"),
    ("Punto Cambio Inmediato", """VAR _l1 = COALESCE([Líneas Acción Cambiar], 0)
VAR _m1 = COALESCE([$ Acción Cambiar], 0)
VAR _l2 = COALESCE([Líneas Acción Solo BO], 0)
VAR _m2 = COALESCE([$ Acción Solo BO], 0)
RETURN
    FORMAT(_l1, "#,0") & " líneas (" & FORMAT(_m1, "$#,0") & ") se pueden cambiar ya al CMMF nuevo, y "
        & FORMAT(_l2, "#,0") & " más (" & FORMAT(_m2, "$#,0") & ") solo en el BO porque el resto ya está en picking.\"""", None, "3. Textos"),
    ("Punto OC Vencida", """VAR _l = COALESCE(CALCULATE([Líneas], KEEPFILTERS('BO Lineas'[Estado OC] = "Vencida")), 0)
VAR _m = COALESCE([$ BO OC Vencida], 0)
RETURN
    IF(_l = 0, "No hay líneas con OC vencida.",
        FORMAT(_l, "#,0") & " líneas con OC vencida suman " & FORMAT(_m, "$#,0") & " (" & FORMAT(DIVIDE(_m, [$ BO]), "0%")
        & " del BO): confirmar con el cliente si aún quiere el producto antes de cambiar el código.")""", None, "3. Textos"),
    ("Punto OC por Vencer", """VAR _l = [Líneas OC por Vencer]
VAR _m = COALESCE([$ BO OC por Vencer], 0)
RETURN
    IF(_l = 0, "Ninguna línea tiene OC que venza en los próximos 7 días.",
        FORMAT(_l, "#,0") & " líneas (" & FORMAT(_m, "$#,0") & ") tienen OC que vence en los próximos 7 días: priorizar su cambio.")""", None, "3. Textos"),
    ("Punto Revisar", """VAR _l = COALESCE([Líneas Acción Revisar], 0)
VAR _m = COALESCE([$ Acción Revisar], 0)
RETURN
    IF(_l = 0, "Ninguna OV trae ya el CMMF nuevo: no hay riesgo de duplicar cantidades.",
        FORMAT(_l, "#,0") & " líneas (" & FORMAT(_m, "$#,0") & ") ya traen el CMMF nuevo en la misma OV: revisar para no duplicar la cantidad.")""", None, "3. Textos"),
    ("Punto Clave", """VAR _orden = SELECTEDVALUE('Puntos Clave'[Orden])
RETURN
    IF(ISBLANK([$ BO]), BLANK(),
        SWITCH(_orden,
            1, [Punto Cliente],
            2, [Punto Producto],
            3, [Punto Cambio Inmediato],
            4, [Punto OC Vencida],
            5, [Punto OC por Vencer],
            6, [Punto Revisar]))""", None, "3. Textos"),
]


def tmdl_columnas(columnas, tabla):
    out = []
    for nombre, tipo, fmt, resumen, oculto, orden, carpeta in columnas:
        out.append(f"\tcolumn {q(nombre)}")
        out.append(f"\t\tdataType: {tipo}")
        if fmt:
            out.append(f"\t\tformatString: {fmt}")
        if oculto:
            out.append("\t\tisHidden")
        if carpeta:
            out.append(f"\t\tdisplayFolder: {carpeta}")
        out.append(f"\t\tlineageTag: {gid(tabla + '|col|' + nombre)}")
        out.append(f"\t\tsummarizeBy: {resumen}")
        out.append(f"\t\tsourceColumn: {nombre}")
        if orden:
            out.append(f"\t\tsortByColumn: {q(orden)}")
        out.append("")
        out.append("\t\tannotation SummarizationSetBy = Automatic")
        if tipo == "dateTime":
            out.append("")
            out.append("\t\tannotation UnderlyingDateTimeDataType = Date")
        out.append("")
    return out


def tmdl_particion(nombre, m):
    lineas = [f"\tpartition {q(nombre)} = m", "\t\tmode: import", "\t\tsource ="]
    lineas += ["\t\t\t\t" + l if l.strip() else "" for l in m.split("\n")]
    lineas.append("")
    return lineas


def tmdl_tabla(nombre, columnas, m, medidas=None):
    out = [f"table {q(nombre)}", f"\tlineageTag: {gid('tabla|' + nombre)}", ""]
    for med, expr, fmt, carpeta in medidas or []:
        if "\n" in expr:
            out.append(f"\tmeasure {q(med)} =")
            out += ["\t\t\t" + l for l in expr.split("\n")]
        else:
            out.append(f"\tmeasure {q(med)} = {expr}")
        if fmt:
            out.append(f"\t\tformatString: {fmt}")
        out.append(f"\t\tdisplayFolder: {carpeta}")
        out.append(f"\t\tlineageTag: {gid('medida|' + med)}")
        out.append("")
    out += tmdl_columnas(columnas, nombre)
    out += tmdl_particion(nombre, m)
    out.append("\tannotation PBI_ResultType = Table")
    out.append("")
    return "\n".join(out)


def escribir(ruta, texto):
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    with open(ruta, "w", encoding="utf-8", newline="\n") as f:
        f.write(texto)


def escribir_json(ruta, obj):
    escribir(ruta, json.dumps(obj, ensure_ascii=False, indent=2) + "\n")


def generar_modelo(raiz, fact_cols, fact_b64, dim_cols, dim_b64):
    sm = os.path.join(raiz, f"{NOMBRE}.SemanticModel")
    d = os.path.join(sm, "definition")
    escribir_json(os.path.join(sm, "definition.pbism"), {
        "$schema": SCHEMA + "item/semanticModel/definitionProperties/1.0.0/schema.json",
        "version": "4.0",
        "settings": {},
    })
    escribir_json(os.path.join(sm, ".platform"), {
        "$schema": SCHEMA + "gitIntegration/platformProperties/2.0.0/schema.json",
        "metadata": {"type": "SemanticModel", "displayName": NOMBRE},
        "config": {"version": "2.0", "logicalId": gid("semanticmodel")},
    })
    escribir(os.path.join(d, "database.tmdl"), "database\n\tcompatibilityLevel: 1567\n\n")
    tablas = [FACT, DIM, MED, PUNTOS]
    escribir(os.path.join(d, "model.tmdl"), "\n".join([
        "model Model",
        "\tculture: en-US",
        "\tdefaultPowerBIDataSourceVersion: powerBI_V3",
        "\tsourceQueryCulture: es-MX",
        "\tdataAccessOptions",
        "\t\tlegacyRedirects",
        "\t\treturnErrorValuesAsNull",
        "",
        "annotation __PBI_TimeIntelligenceEnabled = 0",
        "",
        "annotation PBI_QueryOrder = " + json.dumps(["RutaExcel"] + tablas, ensure_ascii=False),
        "",
        "annotation PBI_ProTooling = [\"DevMode\"]",
        "",
    ] + [f"ref table {q(t)}" for t in tablas] + ["", "ref cultureInfo en-US", ""]))
    escribir(os.path.join(d, "cultures", "en-US.tmdl"), "\n".join([
        "cultureInfo en-US", "",
        "\tlinguisticMetadata =",
        "\t\t\t{",
        '\t\t\t  "Version": "1.0.0",',
        '\t\t\t  "Language": "en-US"',
        "\t\t\t}",
        "\t\tcontentType: json", "", ""]))
    escribir(os.path.join(d, "expressions.tmdl"), "\n".join([
        'expression RutaExcel = "" meta [IsParameterQuery=true, Type="Text", IsParameterQueryRequired=false]',
        f"\tlineageTag: {gid('expr|RutaExcel')}",
        "",
        "\tannotation PBI_ResultType = Text",
        "",
        "\tannotation PBI_NavigationStepName = Navigation",
        "", ""]))
    escribir(os.path.join(d, "relationships.tmdl"), "\n".join([
        f"relationship {gid('rel|cmmf')}",
        f"\tfromColumn: {q(FACT)}.{q('CMMF Viejo')}",
        f"\ttoColumn: {q(DIM)}.{q('CMMF Viejo')}",
        "", ""]))
    escribir(os.path.join(d, "tables", f"{FACT}.tmdl"), tmdl_tabla(FACT, FACT_COLUMNAS, m_fact(fact_cols, fact_b64)))
    escribir(os.path.join(d, "tables", f"{DIM}.tmdl"), tmdl_tabla(DIM, DIM_COLUMNAS, m_dim(dim_cols, dim_b64)))
    escribir(os.path.join(d, "tables", f"{PUNTOS}.tmdl"), tmdl_tabla(PUNTOS, PUNTOS_COLUMNAS, m_puntos()))
    escribir(os.path.join(d, "tables", f"{MED}.tmdl"), tmdl_tabla(
        MED, [("Medida", "string", None, "none", True, None, None)], m_medidas(), MEDIDAS))
    escribir_json(os.path.join(sm, "diagramLayout.json"), {
        "version": "1.1.0",
        "diagrams": [{
            "ordinal": 0, "scrollPosition": {"x": 0, "y": 0}, "nodes": [], "name": "All tables",
            "zoomValue": 100, "pinKeyFieldsToTop": False, "showExtraHeaderInfo": False,
            "hideKeyFieldsWhenCollapsed": False, "tablesLocked": False,
        }],
        "selectedDiagram": "All tables",
        "defaultDiagram": "All tables",
    })
    escribir_json(os.path.join(sm, ".pbi", "editorSettings.json"), {
        "version": "1.0", "parallelQueryLoading": True, "typeDetectionEnabled": True,
        "relationshipImportEnabled": False, "shouldNotifyUserOfNameConflictResolution": True,
    })


# ---------------------------------------------------------------------------
# 3. Reporte (PBIR)
# ---------------------------------------------------------------------------

def lit(v):
    if isinstance(v, bool):
        return {"expr": {"Literal": {"Value": "true" if v else "false"}}}
    if isinstance(v, int) and not isinstance(v, bool):
        return {"expr": {"Literal": {"Value": f"{v}D"}}}
    if isinstance(v, float):
        return {"expr": {"Literal": {"Value": f"{v:g}D"}}}
    return {"expr": {"Literal": {"Value": "'" + v.replace("'", "''") + "'"}}}


def entero(v):
    return {"expr": {"Literal": {"Value": f"{v}L"}}}


def color(h):
    return {"solid": {"color": lit(h)}}


def col(campo, tabla=FACT):
    return {"Column": {"Expression": {"SourceRef": {"Entity": tabla}}, "Property": campo}}


def med(campo):
    return {"Measure": {"Expression": {"SourceRef": {"Entity": MED}}, "Property": campo}}


def proj(field, nombre=None):
    clave = "Column" if "Column" in field else "Measure"
    ent = field[clave]["Expression"]["SourceRef"]["Entity"]
    prop = field[clave]["Property"]
    p = {"field": field, "queryRef": f"{ent}.{prop}", "nativeQueryRef": prop}
    if nombre:
        p["displayName"] = nombre
    if clave == "Column":
        p["active"] = True
    return p


def obj(**props):
    return [{"properties": props}]


def contenedor(titulo=None, fondo=True, borde=True, sombra=True, padding=None, encabezado=True):
    vco = {
        "title": obj(show=lit(bool(titulo)), **({
            "text": lit(titulo), "fontColor": color(TEXTO), "fontSize": lit(11),
            "fontFamily": lit("Segoe UI Semibold"), "alignment": lit("left"),
        } if titulo else {})),
        "background": obj(show=lit(fondo), **({"color": color(BLANCO), "transparency": lit(0)} if fondo else {})),
        "border": obj(show=lit(borde), **({"color": color(BORDE), "radius": lit(8)} if borde else {})),
        "dropShadow": obj(show=lit(sombra), **({
            "color": color("#000000"), "position": lit("Outer"), "preset": lit("Custom"),
            "transparency": lit(90), "shadowBlur": lit(12), "shadowDistance": lit(2),
            "shadowSpread": lit(0), "angle": lit(90),
        } if sombra else {})),
        "visualHeader": obj(show=lit(encabezado)),
    }
    if padding is not None:
        t, r, b, l = padding
        vco["padding"] = obj(top=lit(t), right=lit(r), bottom=lit(b), left=lit(l))
    return vco


class Pagina:
    def __init__(self, clave, nombre):
        self.clave = clave
        self.id = hid("pagina|" + clave)
        self.nombre = nombre
        self.visuales = []

    def agregar(self, clave, x, y, w, h, visual, z=None):
        n = len(self.visuales)
        self.visuales.append({
            "$schema": VC_SCHEMA,
            "name": hid(self.clave + "|" + clave),
            "position": {"x": x, "y": y, "z": z if z is not None else 1000 + n * 100,
                         "height": h, "width": w, "tabOrder": n * 100},
            "visual": visual,
        })


def v_forma(relleno, contorno=None):
    return {
        "visualType": "shape",
        "objects": {
            "shape": [{"properties": {"tileShape": lit("rectangle")}, "selector": {"id": "default"}}],
            "fill": [{"properties": {"show": lit(True), "fillColor": color(relleno), "transparency": lit(0)},
                      "selector": {"id": "default"}}],
            "outline": [{"properties": {"show": lit(bool(contorno)), **({"lineColor": color(contorno), "weight": lit(1)} if contorno else {})},
                         "selector": {"id": "default"}}],
        },
        "visualContainerObjects": contenedor(fondo=False, borde=False, sombra=False, encabezado=False),
        "drillFilterOtherVisuals": True,
    }


def v_texto(texto, tam, color_txt, negrita=True, alineacion="center"):
    return {
        "visualType": "textbox",
        "objects": {
            "general": [{"properties": {"paragraphs": [{
                "textRuns": [{"value": texto, "textStyle": {
                    "fontFamily": "Segoe UI", "fontWeight": "bold" if negrita else "normal",
                    "fontSize": f"{tam}pt", "color": color_txt}}],
                "horizontalTextAlignment": alineacion,
            }]}}],
        },
        "visualContainerObjects": contenedor(fondo=False, borde=False, sombra=False, encabezado=False, padding=(0, 0, 0, 0)),
        "drillFilterOtherVisuals": True,
    }


def v_tarjeta(medida, titulo, unidades=None, decimales=None, tam=22):
    etiquetas = {"color": color(ROJO), "fontSize": lit(tam), "bold": lit(True),
                 "fontFamily": lit("Segoe UI")}
    if unidades is not None:
        etiquetas["labelDisplayUnits"] = lit(unidades)
    if decimales is not None:
        etiquetas["labelPrecision"] = entero(decimales)
    vco = contenedor(padding=(4, 8, 30, 8), encabezado=False)
    vco["title"] = obj(show=lit(True), text=lit(titulo), fontColor=color(TEXTO_2), fontSize=lit(9),
                       bold=lit(True), fontFamily=lit("Segoe UI"), alignment=lit("center"))
    return {
        "visualType": "card",
        "query": {"queryState": {"Values": {"projections": [proj(med(medida))]}}},
        "objects": {"labels": obj(**etiquetas), "categoryLabels": obj(show=lit(False))},
        "visualContainerObjects": vco,
        "drillFilterOtherVisuals": True,
    }


def v_texto_medida(medida, tam, color_txt, negrita=False):
    return {
        "visualType": "card",
        "query": {"queryState": {"Values": {"projections": [proj(med(medida))]}}},
        "objects": {
            "labels": obj(color=color(color_txt), fontSize=lit(tam), bold=lit(negrita), fontFamily=lit("Segoe UI")),
            "categoryLabels": obj(show=lit(False)),
            "wordWrap": obj(show=lit(True)),
        },
        "visualContainerObjects": contenedor(fondo=False, borde=False, sombra=False, encabezado=False, padding=(0, 0, 0, 0)),
        "drillFilterOtherVisuals": True,
    }


def v_segmentador(campo, titulo, grupo):
    vco = contenedor(padding=(2, 6, 2, 6), sombra=False, encabezado=False)
    return {
        "visualType": "slicer",
        "query": {"queryState": {"Values": {"projections": [proj(col(campo))]}}},
        "objects": {
            "data": obj(mode=lit("Dropdown")),
            "header": obj(show=lit(True), text=lit(titulo), textSize=lit(8), fontColor=color(TEXTO_2), bold=lit(True)),
            "items": obj(textSize=lit(9), fontColor=color(TEXTO)),
            "selection": obj(selectAllCheckboxEnabled=lit(True), singleSelect=lit(False)),
        },
        "visualContainerObjects": vco,
        "syncGroup": {"groupName": grupo, "fieldChanges": True, "filterChanges": True},
        "drillFilterOtherVisuals": True,
    }


def v_navegador():
    return {
        "visualType": "pageNavigator",
        "objects": {
            "fill": [
                {"properties": {"show": lit(True), "fillColor": color(BLANCO), "transparency": lit(0)}, "selector": {"id": "default"}},
                {"properties": {"fillColor": color(ROJO_SUAVE)}, "selector": {"id": "hover"}},
                {"properties": {"fillColor": color(ROJO)}, "selector": {"id": "selected"}},
            ],
            "text": [
                {"properties": {"fontColor": color(TEXTO), "fontSize": lit(9), "fontFamily": lit("Segoe UI Semibold")}, "selector": {"id": "default"}},
                {"properties": {"fontColor": color(BLANCO), "bold": lit(True)}, "selector": {"id": "selected"}},
            ],
            "outline": [
                {"properties": {"show": lit(True), "lineColor": color(BORDE), "weight": lit(1)}, "selector": {"id": "default"}},
                {"properties": {"lineColor": color(ROJO)}, "selector": {"id": "selected"}},
            ],
        },
        "visualContainerObjects": contenedor(fondo=False, borde=False, sombra=False, encabezado=False),
        "drillFilterOtherVisuals": True,
    }


def ejes(unidades=1000):
    return {
        "categoryAxis": obj(show=lit(True), fontSize=lit(9), labelColor=color(TEXTO), showAxisTitle=lit(False)),
        "valueAxis": obj(show=lit(False), gridlineShow=lit(False)),
        "labels": obj(show=lit(True), color=color(TEXTO), fontSize=lit(9), labelDisplayUnits=lit(unidades),
                      labelPrecision=entero(0)),
    }


def color_serie(campo, valor, h):
    return {"properties": {"fill": color(h)}, "selector": {"data": [{"scopeId": {"Comparison": {
        "ComparisonKind": 0, "Left": col(campo), "Right": {"Literal": {"Value": "'" + valor + "'"}}}}}]}}


COLORES_ESTADO = [("Vencida", ROJO), ("Vence en 7 días", ROJO_MEDIO), ("Vigente", GRIS)]


def v_barras(tipo, categoria, valores, titulo, orden, serie=None, colores=None, unidades=1000,
             leyenda=False, info=None):
    estado = {"Category": {"projections": [proj(categoria)]},
              "Y": {"projections": [proj(v) for v in valores]}}
    if serie is not None:
        estado["Series"] = {"projections": [proj(serie)]}
    if info:
        estado["Tooltips"] = {"projections": [proj(v) for v in info]}
    objetos = ejes(unidades)
    objetos["dataPoint"] = colores or obj(fill=color(ROJO))
    objetos["legend"] = obj(show=lit(leyenda), **({"position": lit("Top"), "fontSize": lit(9),
                                                   "labelColor": color(TEXTO_2)} if leyenda else {}))
    return {
        "visualType": tipo,
        "query": {"queryState": estado, "sortDefinition": {"sort": [orden]}},
        "objects": objetos,
        "visualContainerObjects": contenedor(titulo),
        "drillFilterOtherVisuals": True,
    }


def v_dona(categoria, valor, titulo, colores=None):
    objetos = {
        "labels": obj(show=lit(True), fontSize=lit(9), color=color(TEXTO), labelStyle=lit("Percent of total"),
                      percentageLabelPrecision=entero(0)),
        "legend": obj(show=lit(True), position=lit("Right"), fontSize=lit(9), labelColor=color(TEXTO)),
        "slices": obj(innerRadiusRatio=entero(60)),
    }
    if colores:
        objetos["dataPoint"] = colores
    return {
        "visualType": "donutChart",
        "query": {"queryState": {"Category": {"projections": [proj(categoria)]},
                                 "Y": {"projections": [proj(valor)]}},
                  "sortDefinition": {"sort": [{"field": valor, "direction": "Descending"}]}},
        "objects": objetos,
        "visualContainerObjects": contenedor(titulo),
        "drillFilterOtherVisuals": True,
    }


def v_mapa_arbol(grupo, detalle, valor, titulo):
    return {
        "visualType": "treemap",
        "query": {"queryState": {"Group": {"projections": [proj(grupo)]},
                                 "Details": {"projections": [proj(detalle)]},
                                 "Values": {"projections": [proj(valor)]}},
                  "sortDefinition": {"sort": [{"field": valor, "direction": "Descending"}]}},
        "objects": {
            "labels": obj(show=lit(True), fontSize=lit(9), color=color(BLANCO), labelDisplayUnits=lit(1000),
                          labelPrecision=entero(0)),
            "categoryLabels": obj(show=lit(True), fontSize=lit(9), color=color(BLANCO)),
            "legend": obj(show=lit(True), position=lit("Top"), fontSize=lit(9), labelColor=color(TEXTO_2)),
        },
        "visualContainerObjects": contenedor(titulo),
        "drillFilterOtherVisuals": True,
    }


def v_tabla(columnas, titulo, orden, anchos=None, rojas=(), totales=True, tam=9, ajustar=False):
    """columnas: lista de (field, nombre_visible)."""
    proyecciones = [proj(f, n) for f, n in columnas]
    objetos = {
        "columnHeaders": obj(fontColor=color(ROJO_OSCURO), backColor=color(ROJO_SUAVE), bold=lit(True),
                             fontSize=lit(tam), wordWrap=lit(True)),
        "values": [{"properties": {"fontSize": lit(tam), "fontColor": color(TEXTO), "wordWrap": lit(ajustar),
                                   "backColorPrimary": color(BLANCO), "backColorSecondary": color(FONDO)}}],
        "grid": obj(gridHorizontal=lit(True), gridHorizontalColor=color(BORDE), gridVertical=lit(False),
                    rowPadding=lit(4)),
        "total": obj(totals=lit(totales), **({"bold": lit(True), "fontColor": color(TEXTO),
                                              "backColor": color(ROJO_SUAVE), "label": lit("Total")} if totales else {})),
    }
    for p in proyecciones:
        if p["queryRef"] in rojas:
            objetos["values"].append({"properties": {"fontColor": color(ROJO), "bold": lit(True)},
                                      "selector": {"metadata": p["queryRef"]}})
    if anchos:
        objetos["columnWidth"] = [{"properties": {"value": lit(w)}, "selector": {"metadata": p["queryRef"]}}
                                  for p, w in zip(proyecciones, anchos) if w]
    return {
        "visualType": "tableEx",
        "query": {"queryState": {"Values": {"projections": proyecciones}},
                  "sortDefinition": {"sort": [orden]}},
        "objects": objetos,
        "visualContainerObjects": contenedor(titulo),
        "drillFilterOtherVisuals": True,
    }


ANCHO, ALTO = 1280, 720
M = 16          # margen
G = 12          # separación
Y0 = 150        # inicio del contenido (debajo de encabezado y segmentadores)
H_KPI = 96


def encabezado(p, titulo, segmentadores=None):
    p.agregar("banda", 0, 0, ANCHO, 60, v_forma(ROJO), z=100)
    p.agregar("titulo", 0, 2, ANCHO, 32, v_texto(titulo, 18, BLANCO), z=200)
    p.agregar("subtitulo", 200, 32, ANCHO - 400, 24, v_texto_medida("Subtítulo Reporte", 10, BLANCO), z=300)
    segs = segmentadores or [("Canal", "Canal"), ("KAM", "KAM"), ("Estado OC", "Estado de la OC"),
                             ("Accion", "Acción a tomar")]
    ancho_nav = 430
    disponible = ANCHO - 2 * M - ancho_nav - 10
    w = (disponible - 8 * (len(segs) - 1)) / len(segs)
    for i, (campo, titulo_seg) in enumerate(segs):
        p.agregar("seg|" + campo, M + i * (w + 8), 66, w, 76, v_segmentador(campo, titulo_seg, "sync_" + campo.replace(" ", "_")))
    p.agregar("navegador", ANCHO - M - ancho_nav, 84, ancho_nav, 40, v_navegador())


def kpi(p, clave, x, y, w, h, medida, titulo, sub, unidades=None, decimales=None):
    p.agregar("kpi|" + clave, x, y, w, h, v_tarjeta(medida, titulo, unidades, decimales))
    p.agregar("kpi_barra|" + clave, x, y, w, 4, v_forma(ROJO))
    p.agregar("kpi_sub|" + clave, x + 6, y + h - 32, w - 12, 28, v_texto_medida(sub, 8, TEXTO_2))


def fila_kpis(p, y, h, items):
    n = len(items)
    w = (ANCHO - 2 * M - G * (n - 1)) / n
    for i, it in enumerate(items):
        kpi(p, it[0], M + i * (w + G), y, w, h, *it[1:])


def ordenar(field, desc=True):
    return {"field": field, "direction": "Descending" if desc else "Ascending"}


def construir_paginas():
    paginas = []

    # --- Página 1: Resumen ejecutivo ------------------------------------------------
    p = Pagina("resumen", "Resumen ejecutivo")
    encabezado(p, "BACKORDER EN CMMF VIEJO  →  CAMBIO AL CMMF NUEVO")
    fila_kpis(p, Y0, H_KPI, [
        ("bo", "$ BO", "$ BACKORDER TOTAL", "Sub $ BO", 1000000, 2),
        ("pzas", "Piezas BO", "PIEZAS EN BACKORDER", "Sub Piezas", 1, 0),
        ("clientes", "Clientes", "CLIENTES AFECTADOS", "Sub Clientes", 1, 0),
        ("vencida", "% BO OC Vencida", "% DEL BO CON OC VENCIDA", "Sub OC Vencida", None, None),
        ("inmediato", "$ Cambio Inmediato", "$ LISTO PARA CAMBIAR", "Sub Cambio Inmediato", 1000000, 2),
        ("porvencer", "$ BO OC por Vencer", "$ CON OC POR VENCER (≤7 DÍAS)", "Sub OC por Vencer", 1000, 0),
    ])
    y2, h2 = Y0 + H_KPI + 10, 214
    p.agregar("acciones", M, y2, 400, h2, v_barras(
        "clusteredBarChart", col("Accion"), [med("$ BO")], "$ BO por acción a tomar",
        ordenar(col("Accion"), desc=False), info=[med("Líneas"), med("% del BO")]))
    p.agregar("clientes", M + 400 + G, y2, 436, h2, v_barras(
        "barChart", col("Cliente Corto"), [med("$ BO")], "Clientes · $ BO por estado de la OC",
        ordenar(med("$ BO")), serie=col("Estado OC"),
        colores=[color_serie("Estado OC", v, c) for v, c in COLORES_ESTADO], leyenda=True,
        info=[med("Líneas")]))
    x3 = M + 400 + G + 436 + G
    p.agregar("canal", x3, y2, ANCHO - M - x3, h2, v_dona(col("Canal"), med("$ BO"), "$ BO por canal"))
    y3 = y2 + h2 + G
    h3 = ALTO - M - y3
    p.agregar("puntos", M, y3, 620, h3, v_tabla(
        [(col("Tema", PUNTOS), "Tema"), (med("Punto Clave"), "Punto clave")], "Puntos clave",
        ordenar(col("Tema", PUNTOS), desc=False), anchos=[130, 440], totales=False, ajustar=True))
    p.agregar("productos", M + 620 + G, y3, ANCHO - 2 * M - 620 - G, h3, v_barras(
        "clusteredBarChart", col("Producto Viejo"), [med("$ BO")], "$ BO por producto (CMMF viejo)",
        ordenar(med("$ BO")), info=[med("Piezas BO"), med("% del BO")]))
    paginas.append(p)

    # --- Página 2: Plan de acción ---------------------------------------------------
    p = Pagina("plan", "Plan de acción")
    encabezado(p, "PLAN DE ACCIÓN  ·  ¿QUÉ HACER CON CADA LÍNEA?")
    fila_kpis(p, Y0, H_KPI, [
        ("cambiar", "$ Acción Cambiar", "CAMBIAR AL CMMF NUEVO", "Sub Acción Cambiar", 1000, 0),
        ("solobo", "$ Acción Solo BO", "CAMBIAR SOLO EL BO (RESTO EN PICKING)", "Sub Acción Solo BO", 1000, 0),
        ("revisar", "$ Acción Revisar", "REVISAR: LA OV YA TRAE EL NUEVO", "Sub Acción Revisar", 1000, 0),
        ("confirmar", "$ Acción Confirmar", "CONFIRMAR CON CLIENTE: OC VENCIDA", "Sub Acción Confirmar", 1000, 0),
    ])
    y2 = Y0 + H_KPI + 10
    h2 = ALTO - M - y2
    p.agregar("prioridades", M, y2, 820, h2, v_tabla([
        (col("Cliente Corto"), "Cliente"),
        (col("KAM"), "KAM"),
        (med("$ BO"), "$ BO"),
        (med("$ BO OC Vencida"), "$ OC vencida"),
        (med("Líneas Acción Cambiar"), "Cambiar al nuevo"),
        (med("Líneas Acción Solo BO"), "Solo el BO"),
        (med("Líneas Acción Revisar"), "Revisar OV"),
        (med("Líneas Acción Confirmar"), "Confirmar cliente"),
        (med("% del BO"), "% del BO"),
    ], "Prioridades por cliente · líneas por acción y KAM responsable", ordenar(med("$ BO")),
        anchos=[150, 130, 90, 90, 70, 60, 60, 70, 60], rojas=(f"{MED}.$ BO OC Vencida",)))
    xr = M + 820 + G
    wr = ANCHO - M - xr
    hr = (h2 - G) / 2
    p.agregar("estado", xr, y2, wr, hr, v_dona(
        col("Estado OC"), med("$ BO"), "Urgencia · $ BO por estado de la OC",
        colores=[color_serie("Estado OC", v, c) for v, c in COLORES_ESTADO]))
    p.agregar("antiguedad", xr, y2 + hr + G, wr, hr, v_barras(
        "clusteredColumnChart", col("Antiguedad BO"), [med("$ BO")], "Antigüedad del BO · días desde el pedido",
        ordenar(col("Antiguedad BO"), desc=False), info=[med("Líneas")]))
    paginas.append(p)

    # --- Página 3: Clientes y productos ---------------------------------------------
    p = Pagina("clientes", "Clientes y productos")
    encabezado(p, "CLIENTES Y PRODUCTOS  ·  ¿DÓNDE ESTÁ EL BACKORDER?")
    ya, ha = Y0, 270
    p.agregar("mapa", M, ya, 620, ha, v_mapa_arbol(col("Canal"), col("Cliente Corto"), med("$ BO"),
                                                   "$ BO por canal y cliente"))
    p.agregar("kam", M + 620 + G, ya, ANCHO - 2 * M - 620 - G, ha, v_barras(
        "clusteredBarChart", col("KAM"), [med("$ BO")], "$ BO por KAM responsable", ordenar(med("$ BO")),
        info=[med("Clientes"), med("Líneas")]))
    yb = ya + ha + G
    p.agregar("mapa_cmmf", M, yb, ANCHO - 2 * M, ALTO - M - yb, v_tabla([
        (col("CMMF Viejo"), "CMMF viejo"),
        (col("Producto Viejo"), "Producto viejo"),
        (col("CMMF Nuevo"), "CMMF nuevo"),
        (col("Producto Nuevo"), "Producto nuevo"),
        (med("Clientes"), "Clientes"),
        (med("Líneas"), "Líneas"),
        (med("Piezas BO"), "Piezas BO"),
        (med("$ BO"), "$ BO"),
        (med("$ BO OC Vencida"), "$ OC vencida"),
        (med("Pzas en Picking Viejo"), "Pzas en picking"),
        (med("% del BO"), "% del BO"),
    ], "Mapa de cambio · CMMF viejo → CMMF nuevo", ordenar(med("$ BO")),
        anchos=[100, 230, 100, 230, 70, 60, 80, 100, 100, 90, 70], rojas=(f"{MED}.$ BO OC Vencida",)))
    paginas.append(p)

    # --- Página 4: Detalle por línea ------------------------------------------------
    p = Pagina("detalle", "Detalle por línea")
    encabezado(p, "DETALLE POR LÍNEA SAP  ·  OV / OC / CMMF", segmentadores=[
        ("Canal", "Canal"), ("KAM", "KAM"), ("Estado OC", "Estado de la OC"), ("Accion", "Acción a tomar"),
        ("Cliente Corto", "Cliente"), ("CMMF Viejo", "CMMF viejo")])
    p.agregar("tabla", M, Y0, ANCHO - 2 * M, ALTO - M - Y0, v_tabla([
        (col("Cliente Corto"), "Cliente"),
        (col("KAM"), "KAM"),
        (col("OV"), "OV"),
        (col("Linea"), "Línea"),
        (col("OC"), "OC"),
        (col("Fecha Pedido"), "Fecha pedido"),
        (col("Dias BO"), "Días BO"),
        (col("Fecha Cancelacion OC"), "Cancela OC"),
        (col("Estado OC"), "Estado OC"),
        (col("CMMF Viejo"), "CMMF viejo"),
        (col("Producto Viejo"), "Producto viejo"),
        (col("CMMF Nuevo"), "CMMF nuevo"),
        (med("Piezas BO"), "Pzas BO"),
        (med("$ BO"), "$ BO"),
        (med("Pzas en Picking Viejo"), "Pzas picking"),
        (col("Que Hacer"), "Qué hacer"),
    ], "Líneas a cambiar · ordenadas por $ BO", ordenar(med("$ BO")),
        anchos=[110, 105, 55, 40, 80, 70, 45, 70, 80, 80, 160, 80, 50, 75, 50, 190], tam=8, ajustar=True))
    paginas.append(p)
    return paginas


def tema():
    return {
        "$schema": "https://raw.githubusercontent.com/microsoft/powerbi-desktop-samples/main/Report%20Theme%20JSON%20Schema/reportThemeSchema-2.140.json",
        "name": CUSTOM_THEME,
        "dataColors": [ROJO, ROJO_OSCURO, ROJO_MEDIO, TEXTO_2, GRIS, "#F4B6B6", "#495057", "#D98C8C",
                       "#343A40", "#CED4DA", "#9E2A2A", "#E9ECEF"],
        "foreground": TEXTO,
        "foregroundNeutralSecondary": TEXTO_2,
        "background": BLANCO,
        "backgroundLight": FONDO,
        "tableAccent": ROJO,
        "good": TEXTO_2,
        "neutral": GRIS,
        "bad": ROJO,
        "maximum": ROJO,
        "center": "#F4B6B6",
        "minimum": ROJO_SUAVE,
        "textClasses": {
            "callout": {"fontFace": "Segoe UI Semibold", "fontSize": 22, "color": ROJO},
            "title": {"fontFace": "Segoe UI Semibold", "fontSize": 11, "color": TEXTO},
            "header": {"fontFace": "Segoe UI Semibold", "fontSize": 11, "color": TEXTO},
            "label": {"fontFace": "Segoe UI", "fontSize": 9, "color": TEXTO_2},
        },
        "visualStyles": {
            "page": {"*": {
                "background": [{"color": {"solid": {"color": FONDO}}, "transparency": 0}],
                "outspace": [{"color": {"solid": {"color": FONDO}}, "transparency": 0}],
            }},
        },
    }


def generar_reporte(raiz, paginas):
    rp = os.path.join(raiz, f"{NOMBRE}.Report")
    d = os.path.join(rp, "definition")
    escribir_json(os.path.join(raiz, f"{NOMBRE}.pbip"), {
        "$schema": SCHEMA + "pbip/pbipProperties/1.0.0/schema.json",
        "version": "1.0",
        "artifacts": [{"report": {"path": f"{NOMBRE}.Report"}}],
        "settings": {"enableAutoRecovery": True},
    })
    escribir_json(os.path.join(rp, ".platform"), {
        "$schema": SCHEMA + "gitIntegration/platformProperties/2.0.0/schema.json",
        "metadata": {"type": "Report", "displayName": NOMBRE},
        "config": {"version": "2.0", "logicalId": gid("report")},
    })
    escribir_json(os.path.join(rp, "definition.pbir"), {
        "$schema": SCHEMA + "item/report/definitionProperties/2.0.0/schema.json",
        "version": "4.0",
        "datasetReference": {"byPath": {"path": f"../{NOMBRE}.SemanticModel"}},
    })
    escribir_json(os.path.join(d, "version.json"), {
        "$schema": SCHEMA + "item/report/definition/versionMetadata/1.0.0/schema.json",
        "version": "2.0.0",
    })
    escribir_json(os.path.join(d, "report.json"), {
        "$schema": SCHEMA + "item/report/definition/report/1.2.0/schema.json",
        "themeCollection": {
            "baseTheme": {"name": BASE_THEME, "reportVersionAtImport": THEME_VERSION, "type": "SharedResources"},
            "customTheme": {"name": CUSTOM_THEME, "reportVersionAtImport": THEME_VERSION, "type": "RegisteredResources"},
        },
        "layoutOptimization": "None",
        "objects": {"section": [{"properties": {"verticalAlignment": lit("Top")}}]},
        "resourcePackages": [
            {"name": "SharedResources", "type": "SharedResources",
             "items": [{"name": BASE_THEME, "path": f"BaseThemes/{BASE_THEME}.json", "type": "BaseTheme"}]},
            {"name": "RegisteredResources", "type": "RegisteredResources",
             "items": [{"name": CUSTOM_THEME, "path": CUSTOM_THEME, "type": "CustomTheme"}]},
        ],
        "settings": {
            "useStylableVisualContainerHeader": True,
            "defaultDrillFilterOtherVisuals": True,
            "allowChangeFilterTypes": True,
            "useDefaultAggregateDisplayName": True,
        },
    })
    escribir_json(os.path.join(rp, "StaticResources", "RegisteredResources", CUSTOM_THEME), tema())
    base = os.path.join(os.path.dirname(os.path.abspath(__file__)), f"{BASE_THEME}.json")
    destino = os.path.join(rp, "StaticResources", "SharedResources", "BaseThemes", f"{BASE_THEME}.json")
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    shutil.copyfile(base, destino)
    escribir_json(os.path.join(d, "pages", "pages.json"), {
        "$schema": SCHEMA + "item/report/definition/pagesMetadata/1.0.0/schema.json",
        "pageOrder": [p.id for p in paginas],
        "activePageName": paginas[0].id,
    })
    for p in paginas:
        escribir_json(os.path.join(d, "pages", p.id, "page.json"), {
            "$schema": SCHEMA + "item/report/definition/page/1.2.0/schema.json",
            "name": p.id,
            "displayName": p.nombre,
            "displayOption": "FitToPage",
            "height": ALTO,
            "width": ANCHO,
            "objects": {
                "background": obj(color=color(FONDO), transparency=lit(0)),
                "outspace": obj(color=color(FONDO), transparency=lit(0)),
            },
        })
        for v in p.visuales:
            escribir_json(os.path.join(d, "pages", p.id, "visuals", v["name"], "visual.json"), v)


def main():
    excel, salida = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(excel, data_only=True)
    fact_cols, fact_filas = leer_hoja(wb["Por linea SAP"], "CANAL")
    dim_cols, dim_filas = leer_hoja(wb["Por CMMF"], "CMMF VIEJO")
    dim_cols, dim_filas = dim_cols[:6], [f[:6] for f in dim_filas]
    faltantes = {a for a, _ in FACT_RENOMBRES} - set(fact_cols)
    if faltantes:
        sys.exit(f"Faltan columnas en 'Por linea SAP': {sorted(faltantes)}")
    for nombre in (f"{NOMBRE}.pbip", f"{NOMBRE}.Report", f"{NOMBRE}.SemanticModel"):
        ruta = os.path.join(salida, nombre)
        if os.path.isdir(ruta):
            shutil.rmtree(ruta)
        elif os.path.exists(ruta):
            os.remove(ruta)
    generar_modelo(salida, fact_cols, embebido(fact_filas), dim_cols, embebido(dim_filas))
    generar_reporte(salida, construir_paginas())
    print(f"{len(fact_filas)} líneas de BO y {len(dim_filas)} CMMF escritos en {salida}")


if __name__ == "__main__":
    main()
