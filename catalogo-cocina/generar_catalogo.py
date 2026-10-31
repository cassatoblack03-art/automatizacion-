#!/usr/bin/env python3
"""Genera el catálogo para clientes (HTML + PDF + PNG) y el Excel interno de precios.

Uso:
    python3 generar_catalogo.py

Datos de entrada: productos.csv (misma carpeta).
  - costo: lo que te cuesta a ti (MXN).
  - aumento_manual: opcional, p. ej. 0.70 = +70 %. Si lo dejas vacío se usa la regla por rango.
  - Filas sin nombre_catalogo o sin costo se mandan a la hoja "Pendientes" y no salen en el catálogo.
"""

import csv
import html
import math
import subprocess
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# CONFIGURACIÓN — cambia aquí lo que quieras
# ---------------------------------------------------------------------------
TITULO = "Catálogo de Cocina"
SUBTITULO = "Licuadoras · Batidoras · Exprimidores · Procesadores"
CONTACTO = ""  # p. ej. "Pedidos por WhatsApp: 55 1234 5678". Vacío = no se muestra.
VIGENCIA = "Precios vigentes: noviembre 2026"

# Reglas de aumento por rango de costo: (costo desde, aumento)
# Entre más barato el producto, mayor % (para que la ganancia en pesos valga la pena);
# entre más caro, menor % (para que el precio siga siendo atractivo).
REGLAS = [
    (0, 0.80),     # costo menor a $300     -> +80 %
    (300, 0.65),   # $300 a $699.99         -> +65 %
    (700, 0.55),   # $700 a $1,199.99       -> +55 %
    (1200, 0.45),  # $1,200 en adelante     -> +45 %
]

ORDEN_CATEGORIAS = [
    "Licuadoras",
    "Licuadoras de alto rendimiento",
    "Licuadoras personales",
    "Batidoras de inmersión",
    "Procesadores y picadores",
    "Exprimidores y extractores",
]

BASE = Path(__file__).resolve().parent
SALIDA = BASE / "salida"
CSV_PATH = BASE / "productos.csv"


# ---------------------------------------------------------------------------
# Lógica de precios (la misma que usan las fórmulas del Excel)
# ---------------------------------------------------------------------------
def aumento_por_regla(costo: float) -> float:
    aplicado = REGLAS[0][1]
    for desde, aumento in REGLAS:
        if costo >= desde:
            aplicado = aumento
    return aplicado


def redondear_precio(bruto: float) -> int:
    """Precios "de tienda": < $1,000 terminan en 9 ($599); >= $1,000 terminan en 49/99 ($1,549)."""
    if bruto < 1000:
        return int(math.ceil(bruto / 10) * 10 - 1)
    return int(math.ceil(bruto / 50) * 50 - 1)


def cargar_productos():
    listos, pendientes = [], []
    with open(CSV_PATH, newline="", encoding="utf-8") as f:
        for fila in csv.DictReader(f):
            costo_txt = (fila.get("costo") or "").strip()
            nombre = (fila.get("nombre_catalogo") or "").strip()
            if not costo_txt or not nombre:
                pendientes.append(fila)
                continue
            costo = float(costo_txt)
            manual_txt = (fila.get("aumento_manual") or "").strip()
            manual = float(manual_txt) if manual_txt else None
            aumento = manual if manual is not None else aumento_por_regla(costo)
            precio = redondear_precio(costo * (1 + aumento))
            listos.append({**fila, "nombre_catalogo": nombre, "costo": costo,
                           "manual": manual, "aumento": aumento, "precio": precio})

    def clave(p):
        cat = p["categoria"]
        idx = ORDEN_CATEGORIAS.index(cat) if cat in ORDEN_CATEGORIAS else len(ORDEN_CATEGORIAS)
        return (idx, cat, p["precio"], p["nombre_catalogo"])

    listos.sort(key=clave)
    return listos, pendientes


def pesos(n: float) -> str:
    return f"${n:,.0f}"


# ---------------------------------------------------------------------------
# Catálogo para clientes (NO muestra costos)
# ---------------------------------------------------------------------------
def generar_html(productos) -> str:
    secciones = []
    categoria_actual = None
    tarjetas = []

    def cerrar_seccion():
        if categoria_actual is not None:
            secciones.append(
                f'<section class="cat"><h2>{html.escape(categoria_actual)}</h2>'
                f'<div class="grid">{"".join(tarjetas)}</div></section>'
            )

    for p in productos:
        if p["categoria"] != categoria_actual:
            cerrar_seccion()
            categoria_actual = p["categoria"]
            tarjetas = []
        desc = (p.get("descripcion") or "").strip()
        tarjetas.append(
            '<article class="card">'
            f'<div class="name">{html.escape(p["nombre_catalogo"])}</div>'
            f'<div class="desc">{html.escape(desc) if desc else "&nbsp;"}</div>'
            f'<div class="price">{pesos(p["precio"])}</div>'
            "</article>"
        )
    cerrar_seccion()

    contacto = f'<div class="contact">{html.escape(CONTACTO)}</div>' if CONTACTO else ""

    return f"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(TITULO)}</title>
<style>
  :root {{
    --ink: #1f2328;
    --muted: #6b7280;
    --line: #e7e2dc;
    --paper: #ffffff;
    --bg: #faf7f3;
    --accent: #c2410c;
    --accent-soft: #fff1e8;
  }}
  * {{ box-sizing: border-box; }}
  html, body {{ margin: 0; background: var(--bg); color: var(--ink);
    font-family: "Inter", "Liberation Sans", Arial, sans-serif; }}
  .page {{ max-width: 820px; margin: 0 auto; padding: 36px 32px 28px; background: var(--paper); }}
  header {{ border-bottom: 3px solid var(--accent); padding-bottom: 16px; margin-bottom: 8px; }}
  h1 {{ margin: 0; font-size: 34px; letter-spacing: -0.5px; }}
  .sub {{ margin-top: 6px; color: var(--muted); font-size: 15px; }}
  .meta {{ margin-top: 10px; display: flex; gap: 10px; flex-wrap: wrap; }}
  .pill {{ background: var(--accent-soft); color: var(--accent); font-weight: 600; font-size: 12px;
    padding: 4px 10px; border-radius: 999px; }}
  .cat {{ margin-top: 22px; break-inside: avoid-page; }}
  h2 {{ font-size: 15px; text-transform: uppercase; letter-spacing: 1.2px; color: var(--accent);
    margin: 0 0 10px; padding-bottom: 6px; border-bottom: 1px solid var(--line); }}
  .grid {{ display: grid; grid-template-columns: repeat(3, 1fr); gap: 10px; }}
  .card {{ border: 1px solid var(--line); border-radius: 10px; padding: 12px 12px 10px;
    display: flex; flex-direction: column; break-inside: avoid; background: #fff; }}
  .name {{ font-weight: 600; font-size: 14px; line-height: 1.25; }}
  .desc {{ color: var(--muted); font-size: 12px; margin-top: 4px; flex: 1; }}
  .price {{ margin-top: 10px; font-size: 22px; font-weight: 800; color: var(--ink); }}
  .price::after {{ content: " MXN"; font-size: 11px; font-weight: 600; color: var(--muted); }}
  footer {{ margin-top: 26px; padding-top: 12px; border-top: 1px solid var(--line);
    color: var(--muted); font-size: 12px; display: flex; justify-content: space-between; gap: 12px;
    flex-wrap: wrap; }}
  .contact {{ color: var(--ink); font-weight: 700; font-size: 14px; }}
  @media (max-width: 640px) {{ .grid {{ grid-template-columns: repeat(2, 1fr); }} .page {{ padding: 24px 16px; }} }}
  @page {{ size: A4; margin: 12mm; }}
  @media print {{ html, body {{ background: #fff; }} .page {{ padding: 0; max-width: none; }} }}
</style>
</head>
<body>
<div class="page">
  <header>
    <h1>{html.escape(TITULO)}</h1>
    <div class="sub">{html.escape(SUBTITULO)}</div>
    <div class="meta"><span class="pill">{len(productos)} productos</span><span class="pill">{html.escape(VIGENCIA)}</span></div>
  </header>
  {"".join(secciones)}
  <footer>
    <div>Precios en pesos mexicanos. Sujetos a disponibilidad.</div>
    {contacto}
  </footer>
</div>
</body>
</html>
"""


def exportar_pdf_png(html_path: Path, pdf_path: Path, png_path: Path):
    script = f"""
const {{ chromium }} = require('playwright');
(async () => {{
  const browser = await chromium.launch();
  const page = await browser.newPage({{ viewport: {{ width: 820, height: 1200 }}, deviceScaleFactor: 2 }});
  await page.goto('file://{html_path}');
  await page.pdf({{ path: '{pdf_path}', format: 'A4', printBackground: true }});
  await page.screenshot({{ path: '{png_path}', fullPage: true }});
  await browser.close();
}})();
"""
    subprocess.run(["node", "-e", script], check=True)


# ---------------------------------------------------------------------------
# Excel interno (con costos y ganancias, fórmulas vivas)
# ---------------------------------------------------------------------------
FUENTE = "Arial"
AZUL = Font(name=FUENTE, color="0000FF")
NEGRO = Font(name=FUENTE)
NEGRITA = Font(name=FUENTE, bold=True)
VERDE = Font(name=FUENTE, color="008000")
TITULO_F = Font(name=FUENTE, bold=True, size=14)
ENCABEZADO_F = Font(name=FUENTE, bold=True, color="FFFFFF")
ENCABEZADO_FILL = PatternFill("solid", fgColor="C2410C")
AMARILLO = PatternFill("solid", fgColor="FFFF00")
GRIS = PatternFill("solid", fgColor="F3F4F6")
BORDE = Border(bottom=Side(style="thin", color="D1D5DB"))
MONEDA = '"$"#,##0.00;("$"#,##0.00);"-"'
MONEDA_ENTERA = '"$"#,##0;("$"#,##0);"-"'
PORCENTAJE = '0.0%;(0.0%);"-"'


def generar_excel(productos, pendientes, ruta: Path):
    wb = Workbook()

    # ---- Hoja Reglas ----
    reglas = wb.active
    reglas.title = "Reglas"
    reglas["A1"] = "Reglas de aumento por rango de costo"
    reglas["A1"].font = TITULO_F
    reglas["A2"] = "Celdas azules = puedes cambiarlas. El Excel recalcula todos los precios solo."
    reglas["A2"].font = Font(name=FUENTE, italic=True, color="6B7280")
    for col, txt in enumerate(["Costo desde (MXN)", "Rango", "Aumento %"], start=1):
        c = reglas.cell(row=4, column=col, value=txt)
        c.font, c.fill = ENCABEZADO_F, ENCABEZADO_FILL
    primera_regla = 5
    for i, (desde, aumento) in enumerate(REGLAS):
        r = primera_regla + i
        hasta = REGLAS[i + 1][0] if i + 1 < len(REGLAS) else None
        reglas.cell(row=r, column=1, value=desde).font = AZUL
        reglas.cell(row=r, column=1).number_format = MONEDA_ENTERA
        rango = f"${desde:,} a ${hasta - 0.01:,.2f}" if hasta else f"${desde:,} en adelante"
        reglas.cell(row=r, column=2, value=rango).font = NEGRO
        reglas.cell(row=r, column=3, value=aumento).font = AZUL
        reglas.cell(row=r, column=3).number_format = PORCENTAJE
    ultima_regla = primera_regla + len(REGLAS) - 1
    nota = ultima_regla + 2
    notas = [
        "Redondeo: si el precio queda abajo de $1,000 se sube al siguiente número terminado en 9 (ej. $593 → $599).",
        "Si queda en $1,000 o más se sube al siguiente terminado en 49 o 99 (ej. $1,473 → $1,499).",
        "Criterio: % más alto en productos baratos (para que la ganancia en pesos valga la pena) y más bajo en los caros",
        "(para que el precio siga siendo atractivo frente a tienda). Supuestos propuestos por Claude; ajústalos a tu gusto.",
        "La columna 'Rango' es solo texto descriptivo: si cambias 'Costo desde', actualízala a mano.",
    ]
    for i, t in enumerate(notas):
        reglas.cell(row=nota + i, column=1, value=t).font = Font(name=FUENTE, color="374151")
    reglas.column_dimensions["A"].width = 20
    reglas.column_dimensions["B"].width = 26
    reglas.column_dimensions["C"].width = 14

    rango_desde = f"Reglas!$A${primera_regla}:$A${ultima_regla}"
    rango_aum = f"Reglas!$C${primera_regla}:$C${ultima_regla}"

    # ---- Hoja Precios ----
    ws = wb.create_sheet("Precios", 0)
    ws["A1"] = "Lista de precios interna — Food Prep (NO compartir con clientes)"
    ws["A1"].font = TITULO_F
    ws["A2"] = ("Azul = dato que puedes cambiar (costo). Amarillo = aumento manual opcional "
                "(ej. 0.70 = +70 %); si lo dejas vacío se usa la regla de la hoja 'Reglas'. "
                "Verde = viene de la hoja 'Reglas'. Negro = fórmula.")
    ws["A2"].font = Font(name=FUENTE, italic=True, color="6B7280")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:L2")
    ws.row_dimensions[2].height = 30

    encabezados = [
        "#", "Categoría", "Producto (como sale en la app)", "Nombre en catálogo", "Costo (MXN)",
        "Aumento por regla", "Aumento manual (opcional)", "Aumento aplicado",
        "Precio sin redondear (MXN)", "Precio de venta (MXN)", "Ganancia por pieza (MXN)", "Margen sobre venta",
    ]
    fila_enc = 4
    for col, txt in enumerate(encabezados, start=1):
        c = ws.cell(row=fila_enc, column=col, value=txt)
        c.font, c.fill = ENCABEZADO_F, ENCABEZADO_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[fila_enc].height = 32

    primera = fila_enc + 1
    for i, p in enumerate(productos):
        r = primera + i
        ws.cell(row=r, column=1, value=i + 1).font = NEGRO
        ws.cell(row=r, column=2, value=p["categoria"]).font = NEGRO
        ws.cell(row=r, column=3, value=p["nombre_app"]).font = NEGRO
        ws.cell(row=r, column=4, value=p["nombre_catalogo"]).font = NEGRO
        c = ws.cell(row=r, column=5, value=p["costo"])
        c.font, c.number_format = AZUL, MONEDA
        c = ws.cell(row=r, column=6, value=f"=INDEX({rango_aum},MATCH(E{r},{rango_desde},1))")
        c.font, c.number_format = VERDE, PORCENTAJE
        c = ws.cell(row=r, column=7, value=p["manual"])
        c.font, c.fill, c.number_format = AZUL, AMARILLO, PORCENTAJE
        c = ws.cell(row=r, column=8, value=f'=IF(G{r}="",F{r},G{r})')
        c.font, c.number_format = NEGRO, PORCENTAJE
        c = ws.cell(row=r, column=9, value=f"=E{r}*(1+H{r})")
        c.font, c.number_format = NEGRO, MONEDA
        c = ws.cell(row=r, column=10, value=f"=IF(I{r}<1000,CEILING(I{r},10)-1,CEILING(I{r},50)-1)")
        c.font, c.number_format = NEGRITA, MONEDA_ENTERA
        c = ws.cell(row=r, column=11, value=f"=J{r}-E{r}")
        c.font, c.number_format = NEGRO, MONEDA
        c = ws.cell(row=r, column=12, value=f'=IF(J{r}=0,"",K{r}/J{r})')
        c.font, c.number_format = NEGRO, PORCENTAJE
        for col in range(1, 13):
            ws.cell(row=r, column=col).border = BORDE
    ultima = primera + len(productos) - 1

    # Totales (1 pieza de cada producto)
    t = ultima + 1
    ws.cell(row=t, column=4, value="TOTAL (1 pieza de cada uno)").font = NEGRITA
    for col, fmt in ((5, MONEDA), (10, MONEDA_ENTERA), (11, MONEDA)):
        L = get_column_letter(col)
        c = ws.cell(row=t, column=col, value=f"=SUM({L}{primera}:{L}{ultima})")
        c.font, c.number_format = NEGRITA, fmt
    c = ws.cell(row=t, column=12, value=f'=IF(J{t}=0,"",K{t}/J{t})')
    c.font, c.number_format = NEGRITA, PORCENTAJE
    for col in range(1, 13):
        ws.cell(row=t, column=col).fill = GRIS
    ws.cell(row=t, column=5).comment = Comment(
        "Costos tomados de las capturas de la app 'Venta Empleado' que mandó el usuario (pestaña Food Prep).",
        "Claude")

    anchos = [5, 26, 36, 40, 13, 12, 14, 12, 15, 14, 15, 12]
    for i, w in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=primera, column=5)

    # ---- Hoja Pendientes ----
    pend = wb.create_sheet("Pendientes")
    pend["A1"] = "Productos con datos incompletos (no salen en el catálogo)"
    pend["A1"].font = TITULO_F
    for col, txt in enumerate(["Producto (app)", "Costo visto (MXN)", "Qué falta"], start=1):
        c = pend.cell(row=3, column=col, value=txt)
        c.font, c.fill = ENCABEZADO_F, ENCABEZADO_FILL
    for i, f in enumerate(pendientes):
        r = 4 + i
        costo = (f.get("costo") or "").strip()
        falta = "Precio (no se ve en la foto)" if not costo else "Nombre (cortado en la foto)"
        pend.cell(row=r, column=1, value=f.get("nombre_app")).font = NEGRO
        c = pend.cell(row=r, column=2, value=float(costo) if costo else None)
        c.font, c.number_format = AZUL, MONEDA
        pend.cell(row=r, column=3, value=falta).font = NEGRO
    r = 4 + len(pendientes) + 1
    pend.cell(row=r, column=1, value=(
        "La app dice 'Artículos encontrados: 37' en Food Prep; en las fotos se alcanzan a ver "
        f"{len(productos) + len(pendientes)}. Manda capturas de los que faltan para completarlo.")).font = \
        Font(name=FUENTE, italic=True, color="6B7280")
    pend.column_dimensions["A"].width = 34
    pend.column_dimensions["B"].width = 18
    pend.column_dimensions["C"].width = 32

    wb.save(ruta)


def main():
    SALIDA.mkdir(exist_ok=True)
    productos, pendientes = cargar_productos()

    html_path = SALIDA / "catalogo_food_prep.html"
    html_path.write_text(generar_html(productos), encoding="utf-8")
    exportar_pdf_png(html_path, SALIDA / "catalogo_food_prep.pdf", SALIDA / "catalogo_food_prep.png")
    generar_excel(productos, pendientes, SALIDA / "precios_internos_food_prep.xlsx")

    print(f"{'Producto':45} {'Costo':>10} {'Aum.':>6} {'Venta':>8} {'Ganancia':>9}")
    for p in productos:
        print(f"{p['nombre_catalogo'][:45]:45} {p['costo']:>10,.2f} {p['aumento']:>6.0%} "
              f"{p['precio']:>8,} {p['precio'] - p['costo']:>9,.2f}")
    tc = sum(p["costo"] for p in productos)
    tv = sum(p["precio"] for p in productos)
    print(f"\n{len(productos)} productos | costo total {tc:,.2f} | venta total {tv:,} | "
          f"ganancia {tv - tc:,.2f} ({(tv - tc) / tv:.1%} margen)")
    print(f"Pendientes: {len(pendientes)}")


if __name__ == "__main__":
    main()
