"""Pruebas de cfdi_a_excel.py. Ejecutar desde la carpeta cfdi-a-excel:

    python -m unittest discover -s tests -v
"""

import shutil
import sys
import tempfile
import unittest
import zipfile
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import cfdi_a_excel as c  # noqa: E402

EJEMPLOS = Path(__file__).resolve().parent.parent / "ejemplos"
XML = EJEMPLOS / "xml"


def leer(nombre):
    return c.leer_cfdi((XML / nombre).read_bytes(), nombre)


class TestLectura(unittest.TestCase):
    def test_cfdi_40(self):
        f = leer("factura_A1001.xml")
        self.assertEqual(f["uuid"], "6F1C2A4B-8D3E-4F5A-9B7C-1D2E3F4A5B6C")
        self.assertEqual((f["serie"], f["folio"], f["version"]), ("A", "1001", "4.0"))
        self.assertEqual(f["rfc_emisor"], "EKU9003173C9")
        self.assertEqual(f["subtotal"], Decimal("1000.00"))
        # El IVA se toma del nodo general, no se suma el de cada concepto
        self.assertEqual(f["iva_trasladado"], Decimal("160.00"))
        self.assertEqual(f["total"], Decimal("1160.00"))
        self.assertEqual(f["forma_pago"], "03 - Transferencia")
        self.assertEqual(f["num_conceptos"], 2)
        self.assertEqual(f["obs"], [])

    def test_cfdi_33(self):
        f = leer("factura_B55_cfdi33.xml")
        self.assertEqual((f["version"], f["serie"], f["folio"]), ("3.3", "B", "55"))
        self.assertEqual(f["total"], Decimal("580.00"))
        self.assertEqual(f["obs"], [])

    def test_retenciones(self):
        f = leer("factura_A1002_honorarios.xml")
        self.assertEqual(f["isr_retenido"], Decimal("1000.00"))
        self.assertEqual(f["iva_retenido"], Decimal("1066.67"))
        self.assertEqual(f["metodo_pago"], "PPD - Pago en parcialidades o diferido")
        self.assertEqual(f["obs"], [])

    def test_complemento_de_pago(self):
        f = leer("pago_P20.xml")
        self.assertEqual(f["tipo"], "P - Pago")
        self.assertEqual(f["monto_pagos"], Decimal("9533.33"))
        self.assertEqual(f["docs_pagados"], ["A2B3C4D5-E6F7-4890-ABCD-EF0123456789"])

    def test_total_que_no_cuadra(self):
        f = leer("factura_C7_total_mal.xml")
        self.assertEqual(len(f["obs"]), 1)
        self.assertIn("no cuadra", f["obs"][0])

    def test_no_es_cfdi(self):
        with self.assertRaisesRegex(ValueError, "No es un CFDI"):
            leer("no_es_cfdi.xml")

    def test_xml_danado(self):
        with self.assertRaisesRegex(ValueError, "dañado"):
            c.leer_cfdi(b"<cfdi:Comprobante Version='4.0'", "roto.xml")

    def test_xml_con_bom_y_acentos_en_ansi(self):
        contenido = (XML / "factura_B55_cfdi33.xml").read_text(encoding="utf-8").encode("cp1252")
        f = c.leer_cfdi(b"\xef\xbb\xbf" + contenido, "ansi.xml")
        self.assertEqual(f["nombre_receptor"], "UNIVERSIDAD ROBOTICA ESPAÑOLA")


class TestCarpeta(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp)

    def test_zip_y_uuid_repetido(self):
        shutil.copy(XML / "factura_A1001.xml", self.tmp)
        with zipfile.ZipFile(self.tmp / "correo.zip", "w") as zf:
            zf.write(XML / "factura_A1001.xml", "copia/factura_A1001.xml")
            zf.writestr("factura.pdf", b"%PDF")
        facturas, errores = c.procesar_carpeta(self.tmp)
        self.assertEqual(errores, [])
        self.assertEqual(len(facturas), 2)
        for f in facturas:
            self.assertIn("UUID repetido", f["observaciones"])

    def test_zip_danado(self):
        (self.tmp / "malo.zip").write_bytes(b"esto no es un zip")
        facturas, errores = c.procesar_carpeta(self.tmp)
        self.assertEqual(facturas, [])
        self.assertEqual(errores[0][0], "malo.zip")


class TestComparacion(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp)

    def test_excel_de_ejemplo(self):
        facturas, _ = c.procesar_carpeta(XML)
        control = c.leer_control(EJEMPLOS / "control_ejemplo.xlsx")
        self.assertEqual(control["hoja"], "Octubre")
        self.assertEqual(control["fila_encabezados"], 3)
        self.assertEqual(control["columnas"], {"uuid": "Folio Fiscal", "folio": "No. Factura",
                                               "rfc_emisor": "RFC Proveedor", "total": "Importe Total"})
        # La fila de TOTAL no cuenta como factura
        self.assertEqual(len(control["registros"]), 5)

        estados = {r["folio"]: r["estado"] for r in c.comparar(facturas, control)}
        self.assertEqual(estados, {
            "A1001": c.OK,                    # UUID en minúsculas en el Excel
            "B55": c.OK,
            "A1002": c.CON_DIFERENCIAS,       # 9,533.00 en el Excel vs 9,533.33
            "C7": c.REVISAR_XML,              # se empareja con "C-0007" por folio; el XML no cuadra
            "P20": c.NO_ESTA_EN_EXCEL,
            "A1003": c.FALTA_XML,
        })

    def test_columnas_indicadas_a_mano_y_folio_numerico(self):
        libro = Workbook()
        hoja = libro.active
        hoja.append(["Documento", "Monto a pagar", "Proveedor RFC"])
        hoja.append([1001, "$1,160.00", "EKU9003173C9"])
        hoja.append([55, 600, "IIA040805DZ4"])
        ruta = self.tmp / "control.xlsx"
        libro.save(ruta)

        with self.assertRaisesRegex(ValueError, "columna"):
            c.leer_control(ruta)  # sin ayuda no reconoce los encabezados
        control = c.leer_control(ruta, manual={"folio": "Documento", "total": "Monto a pagar",
                                               "rfc": "Proveedor RFC"})
        facturas = [leer("factura_A1001.xml"), leer("factura_B55_cfdi33.xml")]
        resultados = {r["folio"]: r for r in c.comparar(facturas, control)}
        self.assertEqual(resultados["A1001"]["estado"], c.OK)
        self.assertEqual(resultados["B55"]["estado"], c.CON_DIFERENCIAS)
        self.assertIn("Total", resultados["B55"]["diferencias"])

    def test_mismo_folio_de_dos_proveedores(self):
        libro = Workbook()
        hoja = libro.active
        hoja.append(["Folio", "RFC", "Total"])
        hoja.append(["7", "EKU9003173C9", 999])
        hoja.append(["7", "IIA040805DZ4", 120])
        ruta = self.tmp / "control.xlsx"
        libro.save(ruta)
        control = c.leer_control(ruta)
        resultado = c.comparar([leer("factura_C7_total_mal.xml")], control)
        emparejado = next(r for r in resultado if r["archivo"])
        self.assertEqual(emparejado["fila"], 3)  # el renglón con el mismo RFC


class TestReporte(unittest.TestCase):
    def test_genera_excel(self):
        with tempfile.TemporaryDirectory() as tmp:
            salida = Path(tmp) / "reporte.xlsx"
            c.main([str(XML), "--comparar", str(EJEMPLOS / "control_ejemplo.xlsx"), "-o", str(salida)])
            libro = load_workbook(salida)
            self.assertEqual(libro.sheetnames, ["Resumen", "Comparación", "Facturas", "Conceptos", "Errores"])
            self.assertEqual(libro["Facturas"].max_row, 6)
            self.assertEqual(libro["Comparación"]["A2"].value, c.CON_DIFERENCIAS)


if __name__ == "__main__":
    unittest.main()
