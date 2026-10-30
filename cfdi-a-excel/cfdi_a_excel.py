#!/usr/bin/env python3
"""
CFDI a Excel
============

Lee los XML de CFDI (facturas electrónicas del SAT, versiones 3.3 y 4.0) de una
carpeta, arma un reporte en Excel y, si le das tu Excel de control, te dice qué
facturas cuadran, cuáles tienen diferencias y cuáles faltan.

Uso básico (desde la carpeta de este archivo):

    python cfdi_a_excel.py
        Lee los XML de xml_entrada/ y, si hay un .xlsx en mi_excel/, lo compara.

    python cfdi_a_excel.py "C:/Facturas/Octubre" --comparar "C:/Control/octubre.xlsx"
        Usa otras rutas.

    python cfdi_a_excel.py --columna total="Importe a pagar" --columna folio="No. Fact"
        Indica a mano cómo se llaman las columnas de tu Excel si no se detectan solas.

Ejecuta `python cfdi_a_excel.py --help` para ver todas las opciones.
"""

from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Falta instalar openpyxl. Ejecuta:  python -m pip install -r requirements.txt")

BASE = Path(__file__).resolve().parent
CARPETA_XML = BASE / "xml_entrada"
CARPETA_CONTROL = BASE / "mi_excel"
CARPETA_REPORTES = BASE / "reportes"

# Diferencia máxima (en pesos) para considerar que dos importes son iguales.
TOLERANCIA = Decimal("0.01")
CERO = Decimal(0)

# --- Catálogos del SAT (solo los más comunes, para que el reporte se lea fácil) ---

TIPOS = {"I": "Ingreso", "E": "Egreso", "T": "Traslado", "N": "Nómina", "P": "Pago",
         "INGRESO": "Ingreso", "EGRESO": "Egreso", "TRASLADO": "Traslado"}
METODOS = {"PUE": "Pago en una sola exhibición", "PPD": "Pago en parcialidades o diferido"}
FORMAS = {"01": "Efectivo", "02": "Cheque nominativo", "03": "Transferencia",
          "04": "Tarjeta de crédito", "05": "Monedero electrónico", "06": "Dinero electrónico",
          "08": "Vales de despensa", "12": "Dación en pago", "13": "Pago por subrogación",
          "14": "Pago por consignación", "15": "Condonación", "17": "Compensación",
          "23": "Novación", "24": "Confusión", "25": "Remisión de deuda",
          "26": "Prescripción o caducidad", "27": "A satisfacción del acreedor",
          "28": "Tarjeta de débito", "29": "Tarjeta de servicios",
          "30": "Aplicación de anticipos", "31": "Intermediario pagos", "99": "Por definir"}
IMPUESTOS = {"001": "ISR", "002": "IVA", "003": "IEPS", "ISR": "ISR", "IVA": "IVA", "IEPS": "IEPS"}

# Columnas de la hoja "Facturas": (encabezado en Excel, llave del diccionario de la factura)
COLUMNAS_FACTURAS = [
    ("Archivo", "archivo"),
    ("UUID (folio fiscal)", "uuid"),
    ("Serie", "serie"),
    ("Folio", "folio"),
    ("Fecha emisión", "fecha"),
    ("Fecha timbrado", "fecha_timbrado"),
    ("Tipo", "tipo"),
    ("RFC emisor", "rfc_emisor"),
    ("Nombre emisor", "nombre_emisor"),
    ("RFC receptor", "rfc_receptor"),
    ("Nombre receptor", "nombre_receptor"),
    ("Uso CFDI", "uso_cfdi"),
    ("Método de pago", "metodo_pago"),
    ("Forma de pago", "forma_pago"),
    ("Moneda", "moneda"),
    ("Tipo de cambio", "tipo_cambio"),
    ("Subtotal", "subtotal"),
    ("Descuento", "descuento"),
    ("IVA trasladado", "iva_trasladado"),
    ("IEPS trasladado", "ieps_trasladado"),
    ("ISR retenido", "isr_retenido"),
    ("IVA retenido", "iva_retenido"),
    ("Total", "total"),
    ("Monto pagado (complemento de pago)", "monto_pagos"),
    ("Facturas que paga (UUID)", "docs_pagados"),
    ("CFDI relacionados", "relacionados"),
    ("Núm. conceptos", "num_conceptos"),
    ("Versión", "version"),
    ("Observaciones", "observaciones"),
]

COLUMNAS_CONCEPTOS = [
    ("Archivo", "archivo"), ("UUID", "uuid"), ("Serie", "serie"), ("Folio", "folio"),
    ("RFC emisor", "rfc_emisor"), ("Clave prod/serv", "clave"),
    ("No. identificación", "no_identificacion"), ("Cantidad", "cantidad"),
    ("Clave unidad", "clave_unidad"), ("Unidad", "unidad"), ("Descripción", "descripcion"),
    ("Valor unitario", "valor_unitario"), ("Descuento", "descuento"), ("Importe", "importe"),
]

# Cómo suelen llamarse las columnas en un Excel de control. Los encabezados se
# comparan sin acentos, espacios ni signos ("No. Factura" -> "nofactura").
# Cada campo tiene: (nombres exactos, nombres que basta con que estén contenidos).
# El orden importa: un encabezado ya asignado no se vuelve a usar.
CAMPOS_EXCEL = {
    "uuid": (["uuid", "foliofiscal", "uuidfiscal", "foliosat", "uuidfactura", "uuidcfdi"],
             ["uuid", "foliofiscal"]),
    "serie": (["serie"], []),
    "folio": (["folio", "factura", "nofactura", "nodefactura", "numfactura", "numerofactura",
               "numerodefactura", "foliofactura", "nofact", "foliointerno", "facturano"],
              ["folio"]),
    "rfc_emisor": (["rfcemisor", "rfcproveedor", "rfcdelemisor", "rfcdelproveedor"], []),
    "rfc_receptor": (["rfcreceptor", "rfccliente", "rfcdelreceptor", "rfcdelcliente"], []),
    "rfc": (["rfc"], []),
    "subtotal": (["subtotal"], []),
    "iva": (["iva", "ivatrasladado", "importeiva", "iva16"], []),
    "total": (["total", "importetotal", "totalfactura", "montototal", "totalapagar",
               "importe", "monto"],
              ["total"]),
}

OK = "OK"
CON_DIFERENCIAS = "Con diferencias"
REVISAR_XML = "Revisar XML"
NO_ESTA_EN_EXCEL = "No está en tu Excel"
FALTA_XML = "Falta el XML"
ORDEN_ESTADOS = [CON_DIFERENCIAS, REVISAR_XML, NO_ESTA_EN_EXCEL, FALTA_XML, OK]
COLOR_ESTADO = {OK: "C6EFCE", CON_DIFERENCIAS: "FFEB9C", REVISAR_XML: "FFE699",
                NO_ESTA_EN_EXCEL: "FFC7CE", FALTA_XML: "F8CBAD"}
COLOR_OBSERVACION = "FFEB9C"


# ---------------------------------------------------------------------------
# Lectura de los XML
# ---------------------------------------------------------------------------

def _nombre(tag):
    """Quita el namespace: '{http://www.sat.gob.mx/cfd/4}Emisor' -> 'Emisor'."""
    return tag.rsplit("}", 1)[-1]


def _hijos(elem, nombre):
    return [h for h in elem if _nombre(h.tag) == nombre] if elem is not None else []


def _hijo(elem, nombre):
    hijos = _hijos(elem, nombre)
    return hijos[0] if hijos else None


def _descendientes(elem, nombre):
    return [e for e in elem.iter() if _nombre(e.tag) == nombre] if elem is not None else []


def _attr(elem, nombre):
    """Lee un atributo sin importar mayúsculas (los CFDI 3.2 usaban minúsculas)."""
    if elem is None:
        return ""
    if nombre in elem.attrib:
        return elem.attrib[nombre].strip()
    buscado = nombre.lower()
    for llave, valor in elem.attrib.items():
        if _nombre(llave).lower() == buscado:
            return valor.strip()
    return ""


def _dec(valor):
    if valor in (None, ""):
        return None
    try:
        return Decimal(str(valor).strip())
    except InvalidOperation:
        return None


def _fecha(valor):
    if not valor:
        return None
    try:
        return datetime.fromisoformat(valor.strip()[:19])
    except ValueError:
        return valor


def _con_descripcion(codigo, catalogo):
    descripcion = catalogo.get(codigo.upper()) if codigo else None
    return f"{codigo} - {descripcion}" if descripcion else codigo


def _tipo(codigo):
    descripcion = TIPOS.get(codigo.upper())
    if not descripcion:
        return codigo
    return f"{codigo.upper()} - {descripcion}" if len(codigo) == 1 else descripcion


def _parsear(contenido):
    contenido = contenido.lstrip(b"\xef\xbb\xbf").lstrip()
    try:
        return ET.fromstring(contenido)
    except ET.ParseError as error:
        # Algunos sistemas dicen "UTF-8" pero guardan en ANSI/Latin-1; se intenta de nuevo.
        try:
            texto = contenido.decode("cp1252", errors="replace")
            texto = re.sub(r"^<\?xml[^>]*\?>", "", texto)
            return ET.fromstring(texto)
        except ET.ParseError:
            raise ValueError(f"El XML está dañado o incompleto ({error})") from None


def leer_cfdi(contenido: bytes, archivo: str) -> dict:
    """Extrae los datos principales de un CFDI. Lanza ValueError si no es un CFDI válido."""
    raiz = _parsear(contenido)
    comp = raiz if _nombre(raiz.tag) == "Comprobante" else next(iter(_descendientes(raiz, "Comprobante")), None)
    if comp is None:
        raise ValueError(f"No es un CFDI: no tiene el nodo Comprobante (el XML empieza con <{_nombre(raiz.tag)}>)")

    emisor = _hijo(comp, "Emisor")
    receptor = _hijo(comp, "Receptor")
    timbre = next(iter(_descendientes(raiz, "TimbreFiscalDigital")), None)

    # Impuestos a nivel comprobante (no los de cada concepto, para no contarlos doble)
    impuestos = _hijo(comp, "Impuestos")
    traslados, retenciones = defaultdict(Decimal), defaultdict(Decimal)
    for nodo in _descendientes(_hijo(impuestos, "Traslados"), "Traslado"):
        clave = _attr(nodo, "Impuesto").upper()
        traslados[IMPUESTOS.get(clave, clave)] += _dec(_attr(nodo, "Importe")) or CERO
    for nodo in _descendientes(_hijo(impuestos, "Retenciones"), "Retencion"):
        clave = _attr(nodo, "Impuesto").upper()
        retenciones[IMPUESTOS.get(clave, clave)] += _dec(_attr(nodo, "Importe")) or CERO
    total_trasladados = _dec(_attr(impuestos, "TotalImpuestosTrasladados"))
    if total_trasladados is None:
        total_trasladados = sum(traslados.values(), CERO)
    total_retenidos = _dec(_attr(impuestos, "TotalImpuestosRetenidos"))
    if total_retenidos is None:
        total_retenidos = sum(retenciones.values(), CERO)

    locales = next(iter(_descendientes(raiz, "ImpuestosLocales")), None)
    locales_trasladados = _dec(_attr(locales, "TotaldeTraslados")) or CERO
    locales_retenidos = _dec(_attr(locales, "TotaldeRetenciones")) or CERO

    # Complemento de pago (REP)
    monto_pagos, docs_pagados = None, []
    pagos = next(iter(_descendientes(raiz, "Pagos")), None)
    if pagos is not None:
        monto_pagos = _dec(_attr(_hijo(pagos, "Totales"), "MontoTotalPagos"))
        if monto_pagos is None:
            monto_pagos = sum((_dec(_attr(p, "Monto")) or CERO for p in _descendientes(pagos, "Pago")), CERO)
        docs_pagados = [_attr(d, "IdDocumento").upper() for d in _descendientes(pagos, "DoctoRelacionado")]

    factura = {
        "archivo": archivo,
        "version": _attr(comp, "Version"),
        "uuid": _attr(timbre, "UUID").upper(),
        "serie": _attr(comp, "Serie"),
        "folio": _attr(comp, "Folio"),
        "fecha": _fecha(_attr(comp, "Fecha")),
        "fecha_timbrado": _fecha(_attr(timbre, "FechaTimbrado")),
        "tipo_codigo": _attr(comp, "TipoDeComprobante").upper(),
        "tipo": _tipo(_attr(comp, "TipoDeComprobante")),
        "rfc_emisor": _attr(emisor, "Rfc").upper(),
        "nombre_emisor": _attr(emisor, "Nombre"),
        "rfc_receptor": _attr(receptor, "Rfc").upper(),
        "nombre_receptor": _attr(receptor, "Nombre"),
        "uso_cfdi": _attr(receptor, "UsoCFDI"),
        "metodo_pago": _con_descripcion(_attr(comp, "MetodoPago"), METODOS),
        "forma_pago": _con_descripcion(_attr(comp, "FormaPago"), FORMAS),
        "moneda": _attr(comp, "Moneda"),
        "tipo_cambio": _dec(_attr(comp, "TipoCambio")),
        "subtotal": _dec(_attr(comp, "SubTotal")),
        "descuento": _dec(_attr(comp, "Descuento")),
        "iva_trasladado": traslados.get("IVA", CERO),
        "ieps_trasladado": traslados.get("IEPS", CERO),
        "isr_retenido": retenciones.get("ISR", CERO),
        "iva_retenido": retenciones.get("IVA", CERO),
        "total": _dec(_attr(comp, "Total")),
        "monto_pagos": monto_pagos,
        "docs_pagados": docs_pagados,
        "relacionados": [_attr(r, "UUID").upper() for r in _descendientes(comp, "CfdiRelacionado")],
        "conceptos": [],
        "obs": [],
    }

    for nodo in _hijos(_hijo(comp, "Conceptos"), "Concepto"):
        factura["conceptos"].append({
            "clave": _attr(nodo, "ClaveProdServ"),
            "no_identificacion": _attr(nodo, "NoIdentificacion"),
            "cantidad": _dec(_attr(nodo, "Cantidad")),
            "clave_unidad": _attr(nodo, "ClaveUnidad"),
            "unidad": _attr(nodo, "Unidad"),
            "descripcion": _attr(nodo, "Descripcion"),
            "valor_unitario": _dec(_attr(nodo, "ValorUnitario")),
            "descuento": _dec(_attr(nodo, "Descuento")),
            "importe": _dec(_attr(nodo, "Importe")),
        })
    factura["num_conceptos"] = len(factura["conceptos"])

    if not factura["uuid"]:
        factura["obs"].append("Sin timbre fiscal (no tiene UUID)")
    subtotal, total = factura["subtotal"], factura["total"]
    if factura["tipo_codigo"][:1] in ("I", "E") and subtotal is not None and total is not None:
        calculado = (subtotal - (factura["descuento"] or CERO) + total_trasladados - total_retenidos
                     + locales_trasladados - locales_retenidos)
        if abs(calculado - total) > TOLERANCIA:
            factura["obs"].append(
                f"El total no cuadra: subtotal - descuento + impuestos = {calculado:,.2f}, "
                f"pero el XML dice {total:,.2f}")
    factura["observaciones"] = "; ".join(factura["obs"])
    return factura


def _archivos_xml(carpeta: Path, errores: list):
    """Devuelve (nombre, contenido) de cada XML de la carpeta, sus subcarpetas y los .zip que contenga."""
    for ruta in sorted(carpeta.rglob("*")):
        if not ruta.is_file():
            continue
        relativo = str(ruta.relative_to(carpeta))
        extension = ruta.suffix.lower()
        if extension == ".xml":
            yield relativo, ruta.read_bytes()
        elif extension == ".zip":
            try:
                with zipfile.ZipFile(ruta) as zf:
                    for info in zf.infolist():
                        if (not info.is_dir() and info.filename.lower().endswith(".xml")
                                and not info.filename.startswith("__MACOSX/")):
                            yield f"{relativo} > {info.filename}", zf.read(info)
            except zipfile.BadZipFile:
                errores.append((relativo, "El .zip está dañado y no se pudo abrir"))


def procesar_carpeta(carpeta: Path):
    """Lee todos los CFDI de la carpeta. Devuelve (facturas, errores)."""
    facturas, errores = [], []
    for nombre, contenido in _archivos_xml(carpeta, errores):
        try:
            facturas.append(leer_cfdi(contenido, nombre))
        except ValueError as error:
            errores.append((nombre, str(error)))

    por_uuid = defaultdict(list)
    for factura in facturas:
        if factura["uuid"]:
            por_uuid[factura["uuid"]].append(factura)
    for repetidas in por_uuid.values():
        if len(repetidas) > 1:
            for factura in repetidas:
                otros = ", ".join(f["archivo"] for f in repetidas if f is not factura)
                factura["obs"].append(f"UUID repetido (también en: {otros})")

    for factura in facturas:
        factura["observaciones"] = "; ".join(factura["obs"])
    return facturas, errores


# ---------------------------------------------------------------------------
# Lectura del Excel de control
# ---------------------------------------------------------------------------

def _normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", texto)


def _clave(valor):
    """Normaliza UUID, folios y RFC para compararlos (sin guiones, espacios ni minúsculas)."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return re.sub(r"[^A-Z0-9Ñ&]", "", str(valor).upper())


def _clave_folio(valor):
    """Como _clave, pero sin ceros a la izquierda en los números: 'A-0012' -> 'A12'."""
    return re.sub(r"(?<![0-9])0+(?=[0-9])", "", _clave(valor))


def _claves_folio_xml(factura):
    return {_clave_folio(factura["folio"]), _clave_folio(factura["serie"] + factura["folio"])} - {""}


def _num_excel(valor):
    if valor is None or valor == "":
        return None
    if isinstance(valor, (int, float)):
        return Decimal(str(valor))
    texto = re.sub(r"[^0-9.\-]", "", str(valor))
    try:
        return Decimal(texto) if texto else None
    except InvalidOperation:
        return None


def _vacio(valor):
    return valor is None or str(valor).strip() == ""


def _texto(valor):
    """Texto de una celda; los números enteros sin '.0' (1001.0 -> '1001')."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _mapear_columnas(encabezados, manual):
    """Decide qué columna del Excel corresponde a cada campo. Devuelve {campo: índice}."""
    normalizados = [_normalizar(e) if e is not None else "" for e in encabezados]
    mapa, usadas = {}, set()
    for campo, nombre in manual.items():
        buscado = _normalizar(nombre)
        if buscado in normalizados:
            mapa[campo] = normalizados.index(buscado)
            usadas.add(mapa[campo])
    for campo, (exactos, contenidos) in CAMPOS_EXCEL.items():
        if campo in mapa or campo in manual:
            continue
        indice = next((normalizados.index(n) for n in exactos
                       if n in normalizados and normalizados.index(n) not in usadas), None)
        if indice is None:
            indice = next((i for n in contenidos for i, enc in enumerate(normalizados)
                           if i not in usadas and n in enc
                           and not (campo == "total" and enc.startswith("sub"))), None)
        if indice is not None:
            mapa[campo] = indice
            usadas.add(indice)
    return mapa


def leer_control(ruta: Path, hoja: str | None = None, manual: dict | None = None) -> dict:
    """Lee tu Excel de control y ubica solo la hoja y la fila de encabezados."""
    manual = manual or {}
    if ruta.suffix.lower() == ".xls":
        raise ValueError(f"'{ruta.name}' es formato .xls antiguo. Ábrelo en Excel y guárdalo como .xlsx")
    if not ruta.is_file():
        raise ValueError(f"No encontré el archivo '{ruta}'")
    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except Exception as error:
        raise ValueError(f"No pude abrir '{ruta.name}': {error}") from None

    if hoja and hoja not in libro.sheetnames:
        raise ValueError(f"Tu Excel no tiene la hoja '{hoja}'. Hojas disponibles: {', '.join(libro.sheetnames)}")

    mejor = None  # (puntaje, hoja, fila_encabezados, mapa, filas)
    for nombre_hoja in [hoja] if hoja else libro.sheetnames:
        filas = [list(f) for f in libro[nombre_hoja].iter_rows(values_only=True)]
        for i, fila in enumerate(filas[:25]):
            mapa = _mapear_columnas(fila, manual)
            if "uuid" not in mapa and "folio" not in mapa:
                continue
            if mejor is None or len(mapa) > mejor[0]:
                mejor = (len(mapa), nombre_hoja, i, mapa, filas)
    libro.close()

    if mejor is None:
        raise ValueError(
            f"No encontré en '{ruta.name}' una columna de UUID/folio fiscal ni de número de factura. "
            "Indícala a mano, por ejemplo:  --columna folio=\"Mi columna\"")
    _, nombre_hoja, fila_encabezados, mapa, filas = mejor
    encabezados = filas[fila_encabezados]
    faltantes = [f"{campo}=\"{nombre}\"" for campo, nombre in manual.items() if campo not in mapa]
    if faltantes:
        disponibles = ", ".join(f'"{e}"' for e in encabezados if not _vacio(e))
        raise ValueError(f"No encontré la columna {', '.join(faltantes)} en la hoja '{nombre_hoja}'. "
                         f"Encabezados disponibles: {disponibles}")

    registros = []
    for numero, fila in enumerate(filas[fila_encabezados + 1:], start=fila_encabezados + 2):
        registro = {campo: fila[i] if i < len(fila) else None for campo, i in mapa.items()}
        if _vacio(registro.get("uuid")) and _vacio(registro.get("folio")):
            continue  # filas vacías, de totales, notas, etc.
        registro["fila"] = numero
        registros.append(registro)

    return {
        "ruta": ruta,
        "hoja": nombre_hoja,
        "fila_encabezados": fila_encabezados + 1,
        "columnas": {campo: str(encabezados[i]) for campo, i in mapa.items()},
        "registros": registros,
    }


# ---------------------------------------------------------------------------
# Comparación XML vs Excel
# ---------------------------------------------------------------------------

def _folio_excel(registro):
    return _clave_folio(_texto(registro.get("serie")) + _texto(registro.get("folio")))


def _monto_xml(factura, campo):
    if campo == "total" and factura["tipo_codigo"] == "P":
        return factura["monto_pagos"]
    return factura[{"subtotal": "subtotal", "iva": "iva_trasladado", "total": "total"}[campo]]


def _diferencias(factura, registro, columnas):
    diferencias = []
    if "uuid" in columnas and not _vacio(registro.get("uuid")) and _clave(registro["uuid"]) != _clave(factura["uuid"]):
        diferencias.append(f"UUID: XML {factura['uuid'] or '(sin UUID)'} / Excel {registro['uuid']}")
    if "folio" in columnas and not _vacio(registro.get("folio")) and _folio_excel(registro) not in _claves_folio_xml(factura):
        diferencias.append(f"Folio: XML {factura['serie']}{factura['folio']} / Excel {_texto(registro['folio'])}")
    for campo, etiqueta, rfcs in (("rfc_emisor", "RFC emisor", [factura["rfc_emisor"]]),
                                  ("rfc_receptor", "RFC receptor", [factura["rfc_receptor"]]),
                                  ("rfc", "RFC", [factura["rfc_emisor"], factura["rfc_receptor"]])):
        if campo in columnas and not _vacio(registro.get(campo)) and _clave(registro[campo]) not in map(_clave, rfcs):
            diferencias.append(f"{etiqueta}: XML {' / '.join(rfcs)} / Excel {registro[campo]}")
    for campo, etiqueta in (("subtotal", "Subtotal"), ("iva", "IVA"), ("total", "Total")):
        if campo not in columnas:
            continue
        en_excel, en_xml = _num_excel(registro.get(campo)), _monto_xml(factura, campo)
        if en_excel is not None and en_xml is not None and abs(en_excel - en_xml) > TOLERANCIA:
            diferencias.append(f"{etiqueta}: XML {en_xml:,.2f} / Excel {en_excel:,.2f} "
                               f"(dif. {en_xml - en_excel:,.2f})")
    return diferencias


def comparar(facturas: list, control: dict) -> list:
    """Empareja cada XML con su renglón del Excel (por UUID y, si no, por folio) y los compara."""
    columnas, registros = control["columnas"], control["registros"]
    usados, pares = set(), {}  # índices de registros ya emparejados / id(factura) -> índice

    por_uuid, por_folio = defaultdict(list), defaultdict(list)
    for i, registro in enumerate(registros):
        if _clave(registro.get("uuid")):
            por_uuid[_clave(registro["uuid"])].append(i)
        if _folio_excel(registro):
            por_folio[_folio_excel(registro)].append(i)

    # 1a pasada: por UUID, que es único
    notas = defaultdict(list)
    for factura in facturas:
        indices = [i for i in por_uuid.get(_clave(factura["uuid"]), []) if i not in usados] if factura["uuid"] else []
        if indices:
            pares[id(factura)] = indices[0]
            usados.update(indices)
            if len(indices) > 1:
                filas = ", ".join(str(registros[i]["fila"]) for i in indices)
                notas[id(factura)].append(f"Aparece {len(indices)} veces en tu Excel (filas {filas})")

    # 2a pasada: por serie/folio (prefiriendo el renglón con el mismo RFC)
    for factura in facturas:
        if id(factura) in pares:
            continue
        candidatos = [i for clave in _claves_folio_xml(factura) for i in por_folio.get(clave, []) if i not in usados]
        if not candidatos:
            continue
        rfcs = {_clave(factura["rfc_emisor"]), _clave(factura["rfc_receptor"])}
        mismo_rfc = [i for i in candidatos
                     if any(_clave(registros[i].get(c)) in rfcs for c in ("rfc_emisor", "rfc_receptor", "rfc"))]
        elegido = (mismo_rfc or candidatos)[0]
        pares[id(factura)] = elegido
        usados.add(elegido)

    resultados = []
    for factura in facturas:
        indice = pares.get(id(factura))
        registro = registros[indice] if indice is not None else None
        if registro is None:
            estado, diferencias = NO_ESTA_EN_EXCEL, []
        else:
            diferencias = _diferencias(factura, registro, columnas) + notas[id(factura)]
            estado = CON_DIFERENCIAS if diferencias else REVISAR_XML if factura["obs"] else OK
        resultados.append({
            "estado": estado,
            "uuid": factura["uuid"],
            "folio": f"{factura['serie']}{factura['folio']}",
            "fecha": factura["fecha"],
            "rfc_emisor": factura["rfc_emisor"],
            "nombre_emisor": factura["nombre_emisor"],
            "total_xml": _monto_xml(factura, "total"),
            "total_excel": _num_excel(registro.get("total")) if registro else None,
            "diferencias": "; ".join(diferencias),
            "obs_xml": factura["observaciones"],
            "archivo": factura["archivo"],
            "fila": registro["fila"] if registro else None,
        })

    for i, registro in enumerate(registros):
        if i not in usados:
            resultados.append({
                "estado": FALTA_XML,
                "uuid": _texto(registro.get("uuid")),
                "folio": _texto(registro.get("serie")) + _texto(registro.get("folio")),
                "fecha": None,
                "rfc_emisor": registro.get("rfc_emisor") or registro.get("rfc"),
                "nombre_emisor": None,
                "total_xml": None,
                "total_excel": _num_excel(registro.get("total")),
                "diferencias": "Está en tu Excel pero no llegó su XML",
                "obs_xml": "",
                "archivo": "",
                "fila": registro["fila"],
            })

    resultados.sort(key=lambda r: ORDEN_ESTADOS.index(r["estado"]))
    return resultados


# ---------------------------------------------------------------------------
# Escritura del reporte
# ---------------------------------------------------------------------------

def _celda(valor):
    if isinstance(valor, Decimal):
        return float(valor)
    if isinstance(valor, list):
        return ", ".join(valor)
    if isinstance(valor, str):
        return ILLEGAL_CHARACTERS_RE.sub("", valor)
    return valor


def _tabla(ws, encabezados, filas, colores=None):
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4E78")
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for i, fila in enumerate(filas):
        ws.append([_celda(v) for v in fila])
        relleno = PatternFill("solid", fgColor=colores[i]) if colores and colores[i] else None
        for celda in ws[ws.max_row]:
            if isinstance(celda.value, float):
                celda.number_format = "#,##0.00"
            elif isinstance(celda.value, datetime):
                celda.number_format = "dd/mm/yyyy hh:mm"
            if relleno:
                celda.fill = relleno
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for columna in ws.iter_cols(min_row=1, max_row=min(ws.max_row, 500)):
        largo = max((17 if isinstance(c.value, datetime) else len(str(c.value)) for c in columna if c.value is not None),
                    default=8)
        ws.column_dimensions[get_column_letter(columna[0].column)].width = min(max(largo + 2, 10), 60)


def escribir_reporte(ruta: Path, carpeta: Path, facturas: list, errores: list,
                     control: dict | None = None, resultados: list | None = None):
    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen"

    if resultados is not None:
        _tabla(libro.create_sheet("Comparación"),
               ["Estado", "Diferencias", "UUID", "Serie y folio", "Fecha", "RFC emisor", "Nombre emisor",
                "Total XML", "Total Excel", "Observaciones del XML", "Archivo XML", "Fila en tu Excel"],
               [[r["estado"], r["diferencias"], r["uuid"], r["folio"], r["fecha"], r["rfc_emisor"],
                 r["nombre_emisor"], r["total_xml"], r["total_excel"], r["obs_xml"], r["archivo"], r["fila"]]
                for r in resultados],
               [COLOR_ESTADO[r["estado"]] for r in resultados])

    _tabla(libro.create_sheet("Facturas"),
           [titulo for titulo, _ in COLUMNAS_FACTURAS],
           [[f.get(llave) for _, llave in COLUMNAS_FACTURAS] for f in facturas],
           [COLOR_OBSERVACION if f["obs"] else None for f in facturas])

    _tabla(libro.create_sheet("Conceptos"),
           [titulo for titulo, _ in COLUMNAS_CONCEPTOS],
           [[c[llave] if llave in c else f[llave] for _, llave in COLUMNAS_CONCEPTOS]
            for f in facturas for c in f["conceptos"]])

    if errores:
        _tabla(libro.create_sheet("Errores"), ["Archivo", "Problema"], [list(e) for e in errores])

    # Hoja de resumen
    lineas = [
        ("Reporte generado", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Carpeta de XML", str(carpeta)),
        ("Facturas leídas", len(facturas)),
        ("Archivos que no se pudieron leer", len(errores)),
        ("Facturas con observaciones (renglones amarillos en 'Facturas')", sum(1 for f in facturas if f["obs"])),
    ]
    totales = defaultdict(Decimal)
    for factura in facturas:
        if factura["total"] is not None and factura["tipo_codigo"] != "P":
            totales[(factura["tipo"], factura["moneda"])] += factura["total"]
    for (tipo, moneda), suma in sorted(totales.items()):
        lineas.append((f"Suma de totales - {tipo} ({moneda})", suma))
    if control is not None:
        lineas += [
            ("", ""),
            ("Excel comparado", str(control["ruta"])),
            ("Hoja", control["hoja"]),
            ("Fila de encabezados", control["fila_encabezados"]),
            ("Columnas usadas", ", ".join(f"{c} = \"{n}\"" for c, n in control["columnas"].items())),
        ]
        for estado in ORDEN_ESTADOS:
            lineas.append((estado, sum(1 for r in resultados if r["estado"] == estado)))
    _tabla(resumen, ["Concepto", "Valor"], [list(l) for l in lineas])
    resumen.auto_filter.ref = None
    if control is not None:
        for fila in resumen.iter_rows(min_row=2):
            if fila[0].value in COLOR_ESTADO:
                fila[0].fill = fila[1].fill = PatternFill("solid", fgColor=COLOR_ESTADO[fila[0].value])

    ruta.parent.mkdir(parents=True, exist_ok=True)
    try:
        libro.save(ruta)
    except PermissionError:
        raise SystemExit(f"No pude guardar '{ruta}'. ¿Lo tienes abierto en Excel? Ciérralo e inténtalo de nuevo.")


# ---------------------------------------------------------------------------
# Programa principal
# ---------------------------------------------------------------------------

def _buscar_control():
    if not CARPETA_CONTROL.is_dir():
        return None
    archivos = [p for p in CARPETA_CONTROL.iterdir() if p.is_file() and not p.name.startswith("~$")]
    for viejo in (p for p in archivos if p.suffix.lower() == ".xls"):
        print(f"  Aviso: '{viejo.name}' es .xls antiguo; guárdalo como .xlsx para poder compararlo.")
    candidatos = sorted((p for p in archivos if p.suffix.lower() in (".xlsx", ".xlsm")),
                        key=lambda p: p.stat().st_mtime, reverse=True)
    if len(candidatos) > 1:
        print(f"  Hay {len(candidatos)} archivos de Excel en mi_excel/; uso el más reciente: {candidatos[0].name}")
    return candidatos[0] if candidatos else None


def _abrir(ruta: Path):
    try:
        if sys.platform.startswith("win"):
            os.startfile(ruta)  # noqa: S606 (solo existe en Windows)
        elif sys.platform == "darwin":
            subprocess.run(["open", str(ruta)], check=False)
        else:
            subprocess.run(["xdg-open", str(ruta)], check=False)
    except Exception as error:
        print(f"  (No pude abrir el archivo automáticamente: {error})")


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")

    parser = argparse.ArgumentParser(
        description="Pasa los XML de CFDI a Excel y los compara contra tu Excel de control.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="Campos para --columna: " + ", ".join(CAMPOS_EXCEL))
    parser.add_argument("carpeta", nargs="?", default=str(CARPETA_XML),
                        help="Carpeta con los XML (también lee subcarpetas y .zip). Por defecto: xml_entrada/")
    parser.add_argument("-o", "--salida", help="Excel a generar. Por defecto: reportes/reporte_cfdi_FECHA.xlsx")
    parser.add_argument("--comparar", metavar="EXCEL",
                        help="Tu Excel de control (.xlsx). Por defecto se usa el que esté en mi_excel/")
    parser.add_argument("--no-comparar", action="store_true", help="Solo extraer los XML, sin comparar")
    parser.add_argument("--hoja", help="Hoja de tu Excel a usar (si no, se detecta sola)")
    parser.add_argument("--columna", action="append", default=[], metavar='CAMPO="Encabezado"',
                        help='Indica a mano una columna de tu Excel, p. ej. --columna total="Monto a pagar"')
    parser.add_argument("--abrir", action="store_true", help="Abrir el reporte al terminar")
    args = parser.parse_args(argv)

    manual = {}
    for texto in args.columna:
        campo, _, encabezado = texto.partition("=")
        campo = campo.strip().lower()
        if campo not in CAMPOS_EXCEL or not encabezado.strip():
            parser.error(f"--columna '{texto}' no es válido. Usa CAMPO=\"Encabezado\" con CAMPO en: "
                         + ", ".join(CAMPOS_EXCEL))
        manual[campo] = encabezado.strip().strip('"')

    carpeta = Path(args.carpeta).expanduser()
    if not carpeta.is_dir():
        if carpeta == CARPETA_XML:
            carpeta.mkdir(parents=True)
        else:
            raise SystemExit(f"No existe la carpeta '{carpeta}'")

    print(f"Leyendo XML de: {carpeta}")
    facturas, errores = procesar_carpeta(carpeta)
    if not facturas and not errores:
        raise SystemExit(f"No hay archivos .xml ni .zip en '{carpeta}'. Copia ahí tus facturas y vuelve a correrlo.")
    print(f"  Facturas leídas: {len(facturas)}  |  Archivos con error: {len(errores)}  |  "
          f"Con observaciones: {sum(1 for f in facturas if f['obs'])}")
    for archivo, problema in errores:
        print(f"    - {archivo}: {problema}")

    control = resultados = None
    ruta_control = None if args.no_comparar else (Path(args.comparar).expanduser() if args.comparar else _buscar_control())
    if ruta_control:
        print(f"Comparando con: {ruta_control}")
        try:
            control = leer_control(ruta_control, args.hoja, manual)
        except ValueError as error:
            raise SystemExit(f"  Error: {error}")
        print(f"  Hoja '{control['hoja']}', encabezados en la fila {control['fila_encabezados']}, "
              f"{len(control['registros'])} renglones con factura")
        print("  Columnas usadas: " + ", ".join(f"{c} = \"{n}\"" for c, n in control["columnas"].items()))
        resultados = comparar(facturas, control)
        print("  " + "  |  ".join(f"{e}: {sum(1 for r in resultados if r['estado'] == e)}" for e in ORDEN_ESTADOS))
    elif not args.no_comparar:
        print("No hay Excel de control en mi_excel/, así que solo se extraen los datos.")

    salida = Path(args.salida) if args.salida else CARPETA_REPORTES / f"reporte_cfdi_{datetime.now():%Y-%m-%d_%H%M%S}.xlsx"
    escribir_reporte(salida, carpeta, facturas, errores, control, resultados)
    print(f"\nReporte guardado en: {salida.resolve()}")
    if args.abrir:
        _abrir(salida)
    return 0


if __name__ == "__main__":
    sys.exit(main())
